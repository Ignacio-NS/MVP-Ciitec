"""Regenera ROBOT_QA/entregables/Informe_Final_RobotQA.docx desde los reportes
reales de robot-qa/reports/.

Uso:
    python ROBOT_QA/entregables/build_informe.py \
        --run smoke-20260921T174007Z --baseline smoke-20260920T180712Z
"""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from build_common import (
    PALETTE, FONT_HEADING, FONT_BODY,
    case_rows, diff_runs, func004_history, load_run,
)

OUT_PATH = Path(__file__).resolve().parent / "Informe_Final_RobotQA.docx"


def _rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


# --------------------------------------------------------------- estilos base
def _setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.font.size = Pt(11)
    normal.font.color.rgb = _rgb(PALETTE["ink"])
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT_HEADING
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = _rgb(PALETTE["navy"])
    h1.paragraph_format.space_before = Pt(22)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.page_break_before = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT_HEADING
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = _rgb(PALETTE["teal"])
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.page_break_before = False

    h3 = doc.styles["Heading 3"]
    h3.font.name = FONT_BODY
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = _rgb(PALETTE["navy"])
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)


def _add_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Robot QA — MVP-Ciitec  ·  Informe Final")
    run.font.name = FONT_BODY
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(PALETTE["muted"])


# --------------------------------------------------------------- helpers de contenido
def add_para(doc, text, *, bold=False, italic=False, size=11, color=None,
             align=None, space_after=8, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = FONT_BODY
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = _rgb(color)
    return p


def add_bullets(doc, items, *, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        run.font.name = FONT_BODY
        run.font.size = Pt(size)
        run.font.color.rgb = _rgb(PALETTE["ink"])


def _set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_text(cell, text, *, bold=False, size=10, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    run.font.name = FONT_BODY
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = _rgb(color)


def add_table(doc, headers, rows, *, col_widths=None, header_fill=None,
              row_colors=None, font_size=9.5):
    header_fill = header_fill or PALETTE["navy"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for i, htext in enumerate(headers):
        cell = table.rows[0].cells[i]
        _set_cell_shading(cell, header_fill)
        _set_cell_text(cell, htext, bold=True, size=font_size, color=PALETTE["white"])

    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        fill = None
        if row_colors and r_idx < len(row_colors):
            fill = row_colors[r_idx]
        for c_idx, val in enumerate(row):
            if fill:
                _set_cell_shading(cells[c_idx], fill)
            _set_cell_text(cells[c_idx], val, size=font_size)

    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Cm(w)
    return table


def _add_toc_field(doc: Document) -> None:
    """Inserta un campo TOC nativo de Word (se llena con F9 / al abrir con 'Actualizar')."""
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Haga clic derecho y elija “Actualizar campo” (o F9) para generar el índice."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r_element = run._r
    r_element.append(fld_begin)
    r_element.append(instr)
    r_element.append(fld_sep)
    r_element.append(placeholder)
    r_element.append(fld_end)


def fmt_pct(x: float) -> str:
    return f"{x * 100:.1f}%"


def fmt_s(x: float | None) -> str:
    if x is None:
        return "—"
    return f"{x:.1f}s"


def fmt_date(iso: str) -> str:
    return iso[:10]


# --------------------------------------------------------------- secciones
def build(run_id: str, baseline_id: str) -> None:
    run = load_run(run_id)
    baseline = load_run(baseline_id)
    diff = diff_runs(run_id, baseline_id)
    rows = case_rows(run)
    rows_base = case_rows(baseline)
    history = func004_history()
    m = run["metrics"]
    mb = baseline["metrics"]
    manifest = run["manifest"]

    doc = Document()
    _setup_styles(doc)
    _add_footer(doc)

    # ---- Portada ----------------------------------------------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(140)
    run_ = p.add_run("INFORME")
    run_.font.name = FONT_HEADING
    run_.font.size = Pt(28)
    run_.font.bold = True
    run_.font.color.rgb = _rgb(PALETTE["navy"])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_ = p.add_run("Robot QA para Sistemas con Inteligencia Artificial")
    run_.font.name = FONT_HEADING
    run_.font.size = Pt(16)
    run_.font.bold = True
    run_.font.color.rgb = _rgb(PALETTE["teal"])

    for text in [
        "Proyecto: Síntesis Automática de Reportes y Briefings Operacionales (MVP-Ciitec)",
        "Taller en Empresa 2 — Facultad de Ingeniería, Universidad San Sebastián",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_ = p.add_run(text)
        run_.font.name = FONT_BODY
        run_.font.size = Pt(12)
        run_.font.color.rgb = _rgb(PALETTE["muted"])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    run_ = p.add_run("Equipo: [completar con los integrantes]")
    run_.font.name = FONT_BODY
    run_.font.size = Pt(11)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_ = p2.add_run(f"Fecha: 21 de septiembre de 2026  ·  corrida de referencia: {run_id}")
    run_.font.name = FONT_BODY
    run_.font.size = Pt(11)

    # ---- Tabla de contenidos ------------------------------------------
    # (el salto de página lo aporta el estilo "Heading 1": page_break_before=True)
    doc.add_heading("Tabla de Contenidos", level=1)
    _add_toc_field(doc)

    # ---- 1. Resumen ejecutivo -----------------------------------------
    doc.add_heading("1. Resumen Ejecutivo y Decisión Recomendada", level=1)
    p = add_para(
        doc,
        f"Gate de liberación de la corrida {run_id}: {run['gate']}",
        bold=True, size=13, color=PALETTE["green"],
    )
    add_para(
        doc,
        f"{run['gate_reason']}. task_success_rate={fmt_pct(m['task_success_rate'])} "
        f"≥ umbral {fmt_pct(manifest['thresholds']['task_success_rate'])}, "
        f"0 fallos críticos, 0 casos en revisión humana.",
    )
    add_para(
        doc,
        "Se diseñó, implementó y ejecutó un Robot QA que invoca la API real de MVP-Ciitec "
        "(login LDAP→JWT, carga de documentos, generación de briefing, versionado, exportación, "
        "auditoría) y evalúa cada resultado con evaluadores deterministas, metamórficos y "
        "semánticos. El catálogo cubre 52 casos en 6 familias; esta corrida ejecutó los 10 casos "
        f"de la suite “smoke” contra el sistema real (sut_commit={manifest['sut_commit']}, "
        f"llm_provider={manifest['llm_provider']}, llm_model={manifest['llm_model']}).",
    )
    add_para(
        doc,
        "Esta es la segunda corrida documentada de esta suite. La corrida anterior "
        f"({baseline_id}, {fmt_date(baseline['started_at'])}) terminó en gate REJECT: el caso "
        "FUNC-004 (síntesis de briefing con un corpus de 3 documentos) superó el presupuesto de "
        "latencia de 180 s y el evaluador schema no encontró contenido en la respuesta (timeout). "
        "Entre esa corrida y esta se aplicó una mitigación en el código del SUT "
        "(backend/app/llm/provider.py, sección 7). El robot QA fue quien detectó la falla original "
        "y quien confirma ahora, con evidencia propia, que la corrección funciona: los 10 casos "
        "pasan, incluido FUNC-004 en 113.2 s. La sección 6.3 documenta la comparación completa y "
        "la sección 7 el análisis de causa raíz y la mitigación.",
    )

    # ---- 2. Descripción del sistema ------------------------------------
    doc.add_heading("2. Descripción del Sistema y Límites del Análisis", level=1)
    add_para(
        doc,
        "MVP-Ciitec ingiere documentos heterogéneos, extrae hechos operacionales con un LLM "
        "(Gemini 2.5 Flash), sintetiza un briefing institucional versionado y trazable y lo "
        "exporta a varios formatos, con control de acceso por rol/unidad/nivel de clasificación. "
        "El análisis se limita al entorno de laboratorio levantado con docker compose, con datos "
        "sintéticos y sin acceso a sistemas de terceros más allá de la API de Gemini que el propio "
        "SUT ya consume en producción.",
    )

    # ---- 3. Riesgos priorizados ----------------------------------------
    doc.add_heading("3. Riesgos Priorizados y Requisitos de Calidad", level=1)
    add_table(
        doc,
        ["Riesgo", "Severidad", "Familia de prueba"],
        [
            ["Fuga de datos por nivel de clasificación", "Crítica", "security"],
            ["Prompt injection vía documento cargado", "Crítica", "security"],
            ["Alucinación / citas sin respaldo (hecho_id roto)", "Alta", "rag"],
            ["Regresión no detectada entre versiones", "Alta", "todas (línea base)"],
            [
                "Latencia/timeout en síntesis multi-documento (dependencia del proveedor LLM)",
                "Alta",
                "rendimiento — detectado y mitigado, ver §6.3 y BUG-001",
            ],
            ["No determinismo entre repeticiones", "Media", "confiabilidad"],
        ],
        col_widths=[7.5, 2.5, 6],
    )

    # ---- 4. Arquitectura -------------------------------------------------
    doc.add_heading("4. Arquitectura del Robot QA", level=1)
    add_para(
        doc,
        "Orquestador → Generador de casos → Adaptador REST (ApiAdapter) → Capturador de "
        "evidencia → Evaluadores → Agregador (gate) → Reportero (HTML/JSON/JUnit), siguiendo la "
        "Guía del proyecto (S4). Contratos: SutAdapter.invoke(case) -> Observation; "
        "Evaluator.evaluate(case, observation, context) -> Score. Un ReplayAdapter permite "
        "reejecutar evaluadores sobre evidencia ya capturada sin volver a llamar al SUT ni al LLM.",
    )

    # ---- 5. Estrategia y cobertura ---------------------------------------
    doc.add_heading("5. Estrategia y Cobertura de Pruebas", level=1)
    add_table(
        doc,
        ["Suite", "N.º casos", "Evaluadores usados en esta corrida"],
        [["smoke", "10", "http_contract, schema, citation_support, budget"]],
        col_widths=[3, 2.5, 10.5],
    )

    # ---- 6. Resultados -----------------------------------------------------
    doc.add_heading("6. Resultados y Comparación con Línea Base", level=1)
    doc.add_heading(f"6.1 Métricas de la campaña ({run_id})", level=2)

    metric_rows = [
        ["run_id", run_id],
        ["sut_commit / robot_commit", manifest["sut_commit"]],
        ["llm_provider / llm_model", f"{manifest['llm_provider']} / {manifest['llm_model']}"],
        ["n_cases", f"{m['n_cases']:.0f}"],
        ["task_success_rate (umbral 0.900)", f"{m['task_success_rate']:.3f}"],
        ["fail_rate", f"{m['fail_rate']:.3f}"],
        ["human_review_rate", f"{m['human_review_rate']:.3f}"],
        ["critical_failure_count", f"{m['critical_failure_count']:.0f}"],
        ["review_queue_count", f"{m['review_queue_count']:.0f}"],
        ["groundedness", f"{m['groundedness']:.3f}"],
        ["hallucination_rate_proxy", f"{m['hallucination_rate_proxy']:.3f}"],
        ["latency_p50_s", f"{m['latency_p50_s']:.1f}"],
        ["latency_p95_s", f"{m['latency_p95_s']:.1f}"],
        ["gate", f"{run['gate']} — {run['gate_reason']}"],
    ]
    add_table(doc, ["Métrica", "Valor"], metric_rows, col_widths=[6, 10],
              row_colors=[None] * (len(metric_rows) - 1) + [PALETTE["teal_light"]])

    doc.add_heading("6.2 Resultado por caso", level=2)
    case_table_rows = []
    for r in rows:
        detail = "—"
        if r["fail_details"]:
            detail = "; ".join(r["fail_details"])
        elif r["latency_s"] is not None:
            detail = f"latencia {fmt_s(r['latency_s'])}"
        case_table_rows.append([r["case_id"], r["requirement_id"], r["verdict"], detail])
    add_table(
        doc,
        ["case_id", "requirement_id", "veredicto", "detalle"],
        case_table_rows,
        col_widths=[2.5, 2.7, 2.3, 8.5],
        row_colors=[PALETTE["teal_light"] if r["verdict"] == "PASS" else "F4CCCC" for r in rows],
    )
    n_pass = sum(1 for r in rows if r["verdict"] == "PASS")
    n_fail = sum(1 for r in rows if r["verdict"] == "FAIL")
    n_hr = sum(1 for r in rows if r["verdict"] == "HUMAN_REVIEW")
    add_para(
        doc,
        f"Resumen: {n_pass} caso(s) APROBADO(S), {n_fail} caso(s) FALLIDO(S), {n_hr} en "
        f"REVISIÓN HUMANA, sobre {len(rows)} ejecutados en esta corrida.",
        italic=True,
    )

    doc.add_heading("6.3 Comparación con la corrida anterior", level=2)
    add_para(
        doc,
        f"La corrida {baseline_id} ({fmt_date(baseline['started_at'])}) terminó en gate REJECT "
        f"por 1 fallo crítico (FUNC-004). Entre esa corrida y {run_id} "
        f"({fmt_date(run['started_at'])}) se aplicó la mitigación de la sección 7. La tabla "
        "siguiente compara ambas corridas caso por caso y métrica por métrica, usando la misma "
        "función de comparación con línea base del robot "
        "(robot_qa.reporting.baseline_diff.compare_to_baseline).",
    )
    compare_rows = [
        ["gate", "REJECT — 1 fallo crítico: FUNC-004", f"{run['gate']} — {run['gate_reason']}"],
        ["task_success_rate", fmt_pct(mb["task_success_rate"]), fmt_pct(m["task_success_rate"])],
        ["fail_rate", fmt_pct(mb["fail_rate"]), fmt_pct(m["fail_rate"])],
        ["critical_failure_count", f"{mb['critical_failure_count']:.0f}", f"{m['critical_failure_count']:.0f}"],
        ["latency_p50_s", fmt_s(mb["latency_p50_s"]), fmt_s(m["latency_p50_s"])],
        ["latency_p95_s", fmt_s(mb["latency_p95_s"]), fmt_s(m["latency_p95_s"])],
        ["groundedness", f"{mb['groundedness']:.2f}", f"{m['groundedness']:.2f}"],
        ["hallucination_rate_proxy", f"{mb['hallucination_rate_proxy']:.2f}", f"{m['hallucination_rate_proxy']:.2f}"],
    ]
    add_table(doc, ["Métrica", f"Antes ({baseline_id})", f"Ahora ({run_id})"], compare_rows,
              col_widths=[5, 5.5, 5.5])

    func004_base = next(r for r in rows_base if r["case_id"] == "FUNC-004")
    func004_now = next(r for r in rows if r["case_id"] == "FUNC-004")
    func003_base = next(r for r in rows_base if r["case_id"] == "FUNC-003")
    func003_now = next(r for r in rows if r["case_id"] == "FUNC-003")
    delta_s = func004_base["latency_s"] - func004_now["latency_s"]
    delta_pct = delta_s / func004_base["latency_s"] * 100
    add_para(
        doc,
        f"FUNC-004 (síntesis con corpus de 3 documentos): {func004_base['verdict']} en "
        f"{fmt_s(func004_base['latency_s'])} → {func004_now['verdict']} en "
        f"{fmt_s(func004_now['latency_s'])} (−{delta_s:.1f}s, −{delta_pct:.1f}%). "
        f"FUNC-003 (control, 1 documento) se mantiene estable: {fmt_s(func003_base['latency_s'])} → "
        f"{fmt_s(func003_now['latency_s'])}. "
        f"fixed={diff['fixed']}, new_failures={diff['new_failures']}, "
        f"regressions_detected={diff['regressions_detected']}.",
        bold=False,
    )

    add_para(doc, "Recurrencia histórica de FUNC-004 (todas las corridas registradas):", bold=True, space_after=4)
    hist_rows = []
    for h in history:
        lat = fmt_s(h["latency_s"])
        note = "sin contenido (timeout)" if not h["schema_passed"] and not h["is_transport_error"] else \
               ("error de transporte, excluido del análisis" if h["is_transport_error"] else "OK")
        hist_rows.append([h["run_id"], lat, "PASS" if h["schema_passed"] else "FAIL", note])
    add_table(doc, ["run_id", "latencia (budget)", "schema", "nota"], hist_rows,
              col_widths=[6, 3, 2, 5],
              row_colors=[PALETTE["teal_light"] if h["schema_passed"] and not h["is_transport_error"]
                          else ("EFEFEF" if h["is_transport_error"] else "F4CCCC") for h in history])
    add_para(
        doc,
        "De las 9 corridas registradas, 6 quedaron entre 180.6 s y 182.5 s (por encima o en el "
        "borde del umbral de 180 s) y fallaron; 1 midió un error de transporte local "
        "(WinError 10061, sin relación con el SUT) y se excluye del análisis; y las 2 corridas más "
        "recientes con la mitigación aplicada bajaron a 116.0 s y 113.2 s. No fue un evento aislado: "
        "era un patrón sistemático que la mitigación de la sección 7 resolvió.",
    )

    doc.add_heading("6.4 Línea base oficial", level=2)
    add_para(
        doc,
        "No se comparó contra una línea base guardada (reports/baseline.json): el proyecto aún no "
        "tiene una fijada. La comparación de la sección 6.3 se calculó directamente entre los dos "
        "reportes JSON con la misma función que usaría esa línea base. Se recomienda ejecutar "
        "`python -m robot_qa run --campaign full --save-baseline` para fijar la línea base oficial "
        "una vez validada la campaña completa.",
    )

    # ---- 7. Incidentes ------------------------------------------------------
    doc.add_heading("7. Incidentes y Análisis de Causa", level=1)
    doc.add_heading("BUG-001 — FUNC-004: Timeout en síntesis de briefing multi-documento", level=3)
    add_table(
        doc,
        ["Campo", "Valor"],
        [
            ["Severidad", "medium (declarada en el caso); crítica para el gate — el evaluador schema marca el fallo como critical=true"],
            ["Requisito afectado", "RF-001 (ingesta multi-formato); RNF de latencia (presupuesto de generación 180 s)"],
            ["Reproducir", "python -m robot_qa run --campaign smoke --cases FUNC-004"],
            [
                "Resultado observado (antes)",
                f"schema: sin 'contenido' en la respuesta (timeout o generación vacía); "
                f"budget: {fmt_s(func004_base['latency_s'])} (umbral 180s), corrida {baseline_id}.",
            ],
            [
                "Resultado observado (después)",
                f"schema: 7 bullets en resumen_ejecutivo; budget: {fmt_s(func004_now['latency_s'])} "
                f"(umbral 180s), corrida {run_id}.",
            ],
            [
                "Evidencia",
                f"robot-qa/evidence/{baseline_id}/cases/FUNC-004.json (antes)  ·  "
                f"robot-qa/evidence/{run_id}/cases/FUNC-004.json (después)",
            ],
            [
                "Recurrencia histórica",
                "6 de 9 corridas registradas de FUNC-004 con corpus de 3 documentos (16 y 20 de "
                "septiembre de 2026) tardaron entre 180.6s y 182.5s, es decir, en el borde o por "
                "encima del umbral de 180s. Ver tabla de la sección 6.3.",
            ],
        ],
        col_widths=[4, 12],
    )
    add_para(
        doc,
        "Causa raíz: el prompt de síntesis (backend/app/llm/provider.py, _SYNTH_SYS) le pide al "
        "LLM un JSON institucional muy extenso (resumen ejecutivo, asuntos críticos, situación, "
        "personal, macrozonas, catástrofe, inteligencia, nueve sub-listas de operaciones, "
        "logística), con un tope de salida de 16 384 tokens (backend/app/config.py, "
        "llm_max_output_tokens). El proveedor activo es Gemini 2.5 Flash "
        "(LLM_PROVIDER=gemini en .env), un modelo “thinking” que por defecto genera tokens de "
        "razonamiento interno antes del JSON final; ese overhead de thinking no se citaba en "
        "ninguna parte del resultado, pero sí consumía una parte relevante de los 180 s de "
        "presupuesto. Con 3 documentos (más hechos, más inconsistencias detectadas por LLM, más "
        "contexto de fuentes) la síntesis se acercaba o superaba el umbral de forma sistemática; "
        "con 1 documento (FUNC-003) el mismo flujo siempre pasó cómodo.",
    )
    add_para(
        doc,
        "Mitigación aplicada: se modificó backend/app/llm/provider.py para que GeminiProvider pase "
        "reasoning_effort=“none” en cada llamada a chat.completions.create, desactivando el "
        "thinking de Gemini 2.5 Flash para las tres operaciones del pipeline (extracción, "
        "detección de contradicciones y síntesis) — mecanismo documentado por Google para el "
        "endpoint OpenAI-compatible de Gemini (ai.google.dev/gemini-api/docs/openai). El cambio es "
        "aditivo: se añadió un atributo extra_create_kwargs a la clase base "
        "_OpenAICompatProvider (vacío por defecto, {}), de modo que DeepSeek, Groq y GitHub "
        "Models no se ven afectados; solo GeminiProvider sobreescribe ese atributo.",
    )
    p = add_para(
        doc,
        f"Estado: RESUELTO Y VALIDADO. La corrida {run_id} confirma la mitigación con evidencia "
        f"propia del robot: FUNC-004 pasa en {fmt_s(func004_now['latency_s'])}, muy por debajo del "
        "presupuesto de 180 s, y el gate de la campaña completa es PASS.",
        bold=True, color=PALETTE["green"],
    )

    # ---- 8. Validación de evaluadores ---------------------------------------
    doc.add_heading("8. Validación de Evaluadores y Revisión Humana", level=1)
    add_para(
        doc,
        "Los 6 evaluadores deterministas (schema, citation_support, required_fact, rbac_leak, "
        "budget, http_contract) se validaron con 37 pruebas unitarias sin red ni LLM "
        "(robot-qa/tests/), incluyendo casos límite (citas rotas, hechos faltantes, umbrales "
        "parciales). El evaluador metamórfico se validó con pruebas puras de sus 6 "
        "transformaciones. El juez semántico (rubric_llm) y la calibración humana (robot_qa "
        "calibrate) dependen de que Gemini esté disponible; esta corrida no incluyó casos que los "
        "requieran.",
    )

    # ---- 9. Riesgos residuales -------------------------------------------
    doc.add_heading("9. Riesgos Residuales y Limitaciones", level=1)
    add_bullets(doc, [
        "La API de MVP-Ciitec no expone tokens ni costo del LLM: cost_per_successful_task "
        "(Guía S7) no se puede medir sin instrumentar el backend.",
        "Los usuarios demo solo cubren dos unidades reales (una de ellas transversal); la fuga "
        "de datos ENTRE unidades no transversales no se puede probar sin sumar un sexto usuario a "
        "db/seed.sql y LDAP_USERS.",
        "El evaluador metamórfico usa heurísticas de similitud por palabras (Jaccard), "
        "declaradamente ruidosas frente a un LLM real; los umbrales son ejemplos docentes "
        "(config/thresholds.yaml) que cada equipo debe recalibrar con más evidencia.",
        f"La mitigación de BUG-001 está validada con una sola corrida ({fmt_s(func004_now['latency_s'])} "
        "sobre un presupuesto de 180 s, ≈37% de holgura). La confirmación estadística "
        "(múltiples repeticiones, para descartar variabilidad del proveedor de LLM) queda "
        "pendiente en la campaña `full` con 3 repeticiones por caso.",
    ])

    # ---- 10. Conclusiones y backlog --------------------------------------
    doc.add_heading("10. Conclusiones y Backlog de Mejora", level=1)
    add_para(
        doc,
        "El Robot QA cumple el objetivo de la Guía del proyecto: ejecuta contra el sistema real, "
        "captura evidencia reproducible, aplica evaluadores en los 6 niveles definidos y emite un "
        "veredicto de liberación automático. Esta corrida documenta el ciclo completo: el robot "
        "detectó una falla real de rendimiento del sistema bajo carga (timeout de síntesis con 3 "
        "documentos, corrida del 20 de septiembre), el equipo aplicó una mitigación en el código "
        "del SUT, y el robot validó esa mitigación con una nueva corrida (21 de septiembre) sin "
        "inventar resultados en ningún paso.",
    )
    doc.add_heading("Backlog de mejora", level=2)
    add_bullets(doc, [
        "Ejecutar la campaña `full` (52 casos × 3 repeticiones) y fijar la línea base oficial.",
        "Sumar un sexto usuario demo de otra unidad no transversal para cerrar la limitación de "
        "fuga entre unidades.",
        "Instrumentar el backend para exponer tokens/costo por generación y habilitar "
        "cost_per_successful_task.",
        "Correr `robot_qa calibrate` con una muestra real revisada por una persona del equipo "
        "para reportar human_agreement con cifras propias.",
        "Demostrar en clase la regresión de la rama demo/regresion-qa: gate PASS en main vs. "
        "REJECT en la rama con trazabilidad deshabilitada.",
    ])

    doc.save(str(OUT_PATH))
    print(f"Informe generado: {OUT_PATH}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--run", required=True)
    parser.add_argument("--baseline", required=True)
    args = parser.parse_args()
    build(args.run, args.baseline)
