"""Utilidades compartidas por build_informe.py y build_presentacion.py.

Lee los reportes JSON reales de robot-qa/reports/ y la evidencia cruda de
robot-qa/evidence/ para que los entregables (.docx / .pptx) se generen a
partir de datos verificables, no de cifras escritas a mano.

Uso como selftest (verifica el diff antes/ahora sin generar documentos):
    python ROBOT_QA/entregables/build_common.py --selftest
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
ROBOT_QA_DIR = REPO_ROOT / "robot-qa"
REPORTS_DIR = ROBOT_QA_DIR / "reports"
EVIDENCE_DIR = ROBOT_QA_DIR / "evidence"
ASSETS_DIR = Path(__file__).resolve().parent / "assets"

# El paquete robot_qa (pydantic models + baseline_diff) vive en robot-qa/src.
sys.path.insert(0, str(ROBOT_QA_DIR / "src"))


# --------------------------------------------------------------- Paleta / tipografía
PALETTE = {
    "navy": "1B3A5C",
    "navy_deep": "12233A",
    "teal": "0D9488",
    "teal_light": "A9DFD8",
    "pale": "CADCFC",
    "ink": "1B2430",
    "muted": "5B6B7A",
    "white": "FFFFFF",
    "red": "C0392B",
    "amber": "B08800",
    "green": "1A7F37",
    "rule": "D8DEE6",
}
FONT_HEADING = "Cambria"
FONT_BODY = "Calibri"
FONT_MONO = "Consolas"


# --------------------------------------------------------------- Carga de reportes
def load_run(run_id: str) -> dict:
    """Lee robot-qa/reports/<run_id>.json (salida real de `robot_qa run`)."""
    path = REPORTS_DIR / f"{run_id}.json"
    if not path.is_file():
        raise FileNotFoundError(f"No existe el reporte: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def case_rows(run: dict) -> list[dict]:
    """Aplana run['cases'] a filas planas y fáciles de tabular.

    Forma real de cada entrada en run['cases']:
        {"case": {...TestCase...}, "observations": [...], "scores": [...],
         "verdict": "PASS" | "FAIL" | ..., "consistency": ..., "notes": ...}
    """
    rows = []
    for entry in run["cases"]:
        case = entry["case"]
        scores = entry.get("scores") or []
        budget_scores = [s["value"] for s in scores if s.get("evaluator_id") == "budget"]
        fail_details = [
            f"{s['evaluator_id']}: {s['detail']}"
            for s in scores
            if not s.get("passed", True)
        ]
        rows.append({
            "case_id": case["case_id"],
            "requirement_id": case["requirement_id"],
            "category": case["category"],
            "title": case.get("title", ""),
            "action": (case.get("input") or {}).get("action", ""),
            "verdict": entry["verdict"],
            "latency_s": budget_scores[0] if budget_scores else None,
            "scores": scores,
            "fail_details": fail_details,
        })
    return rows


# --------------------------------------------------------------- Diff entre corridas
def diff_runs(current_id: str, baseline_id: str) -> dict:
    """Reutiliza robot_qa.reporting.baseline_diff.compare_to_baseline."""
    from robot_qa.models import CampaignResult
    from robot_qa.reporting.baseline_diff import compare_to_baseline

    current = CampaignResult.model_validate(load_run(current_id))
    baseline = CampaignResult.model_validate(load_run(baseline_id))
    return compare_to_baseline(current, baseline)


# --------------------------------------------------------------- Historia de FUNC-004
def func004_history() -> list[dict]:
    """Recorre robot-qa/evidence/*/cases/FUNC-004.json en orden cronológico."""
    history = []
    for evidence_file in sorted(EVIDENCE_DIR.glob("*/cases/FUNC-004.json")):
        run_id = evidence_file.parent.parent.name
        data = json.loads(evidence_file.read_text(encoding="utf-8"))
        scores = data.get("scores") or []
        budget = next((s for s in scores if s.get("evaluator_id") == "budget"), None)
        schema = next((s for s in scores if s.get("evaluator_id") == "schema"), None)
        is_transport_error = bool(schema and "error de transporte" in (schema.get("detail") or ""))
        history.append({
            "run_id": run_id,
            "latency_s": budget["value"] if budget else None,
            "schema_passed": schema["passed"] if schema else None,
            "schema_detail": schema["detail"] if schema else "",
            "is_transport_error": is_transport_error,
        })
    return history


# --------------------------------------------------------------- Iconos del deck
def extract_icons(pptx_path: Path, out_dir: Path) -> list[Path]:
    """Vuelca las imágenes únicas embebidas en un .pptx a out_dir. Idempotente."""
    from pptx import Presentation

    out_dir.mkdir(parents=True, exist_ok=True)
    existing = sorted(out_dir.glob("icon_*.png"))
    if existing:
        return existing

    prs = Presentation(str(pptx_path))
    seen: dict[int, Path] = {}
    saved: list[Path] = []
    i = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:  # PICTURE
                blob = shape.image.blob
                key = hash(blob)
                if key in seen:
                    continue
                i += 1
                fn = out_dir / f"icon_{i}.png"
                fn.write_bytes(blob)
                seen[key] = fn
                saved.append(fn)
    return saved


# --------------------------------------------------------------- Self-test
def _selftest() -> int:
    diff = diff_runs("smoke-20260921T174007Z", "smoke-20260920T180712Z")
    print("fixed =", diff["fixed"])
    print("new_failures =", diff["new_failures"])
    print("regressions_detected =", diff["regressions_detected"])
    for key in ("task_success_rate", "latency_p50_s", "latency_p95_s", "fail_rate", "critical_failure_count"):
        d = diff["metric_diffs"].get(key)
        if d:
            print(f"  {key}: baseline={d['baseline']} current={d['current']} delta={d['delta']:+.3f} regressed={d['regressed']}")

    hist = func004_history()
    print("\nFUNC-004 historia:")
    for h in hist:
        flag = " (error de transporte, excluido)" if h["is_transport_error"] else ""
        lat = f"{h['latency_s']:.1f}s" if h["latency_s"] is not None else "?"
        print(f"  {h['run_id']}: {lat} schema_passed={h['schema_passed']}{flag}")

    ok = diff["fixed"] == ["FUNC-004"] and diff["new_failures"] == [] and not diff["regressions_detected"]
    print("\nSELFTEST", "OK" if ok else "FALLÓ")
    return 0 if ok else 1


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--selftest", action="store_true")
    args = parser.parse_args()
    if args.selftest:
        raise SystemExit(_selftest())
    parser.print_help()
