"""Hojas adicionales configurables: render del reporte, esquema y export Word/texto."""
import io
from datetime import datetime

from pptx import Presentation

from app.pipeline import exportacion, reporte_pptx
from app.schemas.llm import BriefingOut, HojaAdicional, HojaCompletada

FECHA = datetime(2026, 10, 4)


def _n_paginas(contenido: dict) -> int:
    return len(Presentation(io.BytesIO(reporte_pptx.construir_pptx(contenido, FECHA))).slides)


def _titulos(contenido: dict) -> list[str | None]:
    prs = Presentation(io.BytesIO(reporte_pptx.construir_pptx(contenido, FECHA)))
    out = []
    for s in prs.slides:
        t = next((sh for sh in s.shapes if sh.name == "object 16" and sh.has_text_frame), None)
        out.append(t.text_frame.text if t is not None else None)
    return out


def test_sin_hojas_sigue_siendo_6_paginas():
    assert _n_paginas({}) == 6
    assert _n_paginas({"hojas_adicionales": []}) == 6


def test_hoja_al_final_y_tras_pagina():
    c = {"hojas_adicionales": [
        {"id": "a", "titulo": "Final", "despues_de": None,
         "bloques": [{"tipo": "texto", "texto": "hola"}]},
        {"id": "b", "titulo": "Tras portada", "despues_de": 0,
         "bloques": [{"tipo": "vinetas", "items": ["x", "y"]}]},
    ]}
    t = _titulos(c)
    assert len(t) == 8
    assert t[1] == "TRAS PORTADA"   # justo después de la portada
    assert t[-1] == "FINAL"         # al final


def test_tabla_larga_pagina_con_cont():
    filas = [[f"UNIDAD {i}", "detalle " * 5, str(i)] for i in range(80)]
    c = {"hojas_adicionales": [{"id": "a", "titulo": "Tabla", "bloques": [
        {"tipo": "tabla", "titulo": "Unidades", "columnas": ["U", "D", "F"], "filas": filas}]}]}
    t = _titulos(c)
    assert len(t) > 7
    assert t[-1] == "TABLA (CONT.)"


def test_todos_los_tipos_de_bloque_y_bloques_vacios_no_rompen():
    c = {"hojas_adicionales": [{"id": "a", "titulo": "Mix", "bloques": [
        {"tipo": "texto", "titulo": "T", "texto": "uno\n\ndos"},
        {"tipo": "vinetas", "items": ["a", "", "b"]},
        {"tipo": "tabla", "columnas": ["A", "B"], "filas": [["1", "2"], ["3"]]},
        {"tipo": "kpis", "items": [{"etiqueta": "F", "valor": "1.234"}, "suelto"]},
        {"tipo": "grafico", "estilo": "barras_h", "categorias": ["a", "b"], "valores": [3, "4.5"]},
        {"tipo": "grafico", "estilo": "dona", "categorias": ["a"], "valores": [1]},
        {"tipo": "imagen", "objeto": "no-existe"},      # sin MinIO: se omite
        {"tipo": "texto", "texto": ""},                  # vacío: se omite
        {"tipo": "tabla", "columnas": [], "filas": []},  # vacío: se omite
    ]}, {"id": "b", "titulo": "Vacía", "bloques": []}]}
    # 6 de la plantilla + la hoja "Mix" (≥1 página, puede paginar) + la hoja vacía (1 página)
    assert _n_paginas(c) >= 8


def test_texto_gigante_se_reparte_en_paginas():
    c = {"hojas_adicionales": [{"id": "a", "titulo": "Largo", "bloques": [
        {"tipo": "texto", "titulo": "Informe", "texto": "palabra " * 4000}]}]}
    assert _n_paginas(c) > 7


def test_grafico_sin_datos_numericos_se_omite():
    assert reporte_pptx._chart_generico("barras", ["a"], ["n/a"]) is None
    assert reporte_pptx._chart_generico("barras", ["a", "b"], [12.5, 30]) is not None


def test_esquema_es_permisivo_y_normaliza_tipo():
    h = HojaAdicional.model_validate({"id": "x", "bloques": [{"tipo": "raro", "basura": 1}]})
    assert h.bloques[0].tipo == "texto"
    b = BriefingOut.model_validate({"hojas_adicionales": [{"titulo": "T"}]})
    assert b.hojas_adicionales[0].titulo == "T"
    assert BriefingOut.model_validate({}).hojas_adicionales == []
    assert HojaCompletada.model_validate({"bloques": [{"tipo": "kpis"}]}).trazabilidad == {}


def test_word_y_texto_incluyen_las_hojas():
    c = {"hojas_adicionales": [{"seccion": "Inteligencia", "titulo": "Frontera", "bloques": [
        {"tipo": "texto", "titulo": "Resumen", "texto": "contenido clave"},
        {"tipo": "tabla", "columnas": ["A", "B"], "filas": [["1", "2"]]}]}]}
    txt = exportacion.render_texto(c, "t", FECHA).decode()
    assert "FRONTERA" in txt and "contenido clave" in txt and "1 | 2" in txt
    assert len(exportacion.render_word(c, "t", FECHA)) > 1000
