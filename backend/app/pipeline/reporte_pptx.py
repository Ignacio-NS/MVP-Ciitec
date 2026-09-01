"""
Filtro de salida — REPORTE INSTITUCIONAL en PDF (RF-008).

Se RELLENA la plantilla .pptx de 6 páginas (``templates/reporte_template.pptx``) con
los datos del briefing (JSON ``contenido`` == ``BriefingOut``) y se convierte a PDF con
LibreOffice headless. NO se genera el .pptx desde cero: se abre la plantilla y se editan
sus shapes in-place.

Reglas (fijadas con el usuario):
  * Los shapes de la plantilla se llaman genéricamente ``object N``; se mapean por
    ``(slide_idx_0based, shape.name)`` con los bindings de cada función ``_sN_*``.
  * Texto: se edita ``runs[0].text`` conservando el formato. Tablas: se rellenan desde
    la 1ª fila de datos (``start_row``), clonando la última fila (<a:tr>) si faltan.
  * Dato ausente en el JSON -> celda en blanco. NUNCA se inventa (los ejemplos horneados
    en la plantilla se limpian al no haber dato).
  * Los "gráficos" de la plantilla son formas dibujadas (freeform), no charts nativos:
    se regeneran con matplotlib (PNG, dpi=200) y se insertan en las MISMAS coordenadas,
    eliminando la(s) forma(s) original(es). Sin datos -> se omite el gráfico.

LibreOffice NO aplica el autoajuste de PowerPoint ("shrink text on overflow"): si un
texto no cabe al tamaño horneado lo PARTE (incluso dentro de una palabra -> letras en
vertical) y se sale del recuadro. Por eso:
  * las fuentes de la plantilla (Franklin Gothic Medium, Trebuchet MS, Calibri, Times)
    van embebidas en ``templates/fonts/`` y se instalan en la imagen Docker;
  * los anchos se verifican con la MÉTRICA REAL de esas fuentes (PIL) y se fijan
    tamaños/anchos de columna que caben, y
  * las tablas que crecen (filas clonadas / textos largos) reposicionan los bloques
    inferiores (reflow) para que nada se superponga.

Gráficos generados: dona Planta/Variable (portada mini + pág. 2), barras de fuerza por
macrozona (pág. 2), barras de fuerza de catástrofe (pág. 2), traslados por UAC (pág. 5)
y operacionalidad logística Total/NOP (pág. 6).
"""
from __future__ import annotations

import io
import os
import re
from copy import deepcopy
from datetime import datetime
from functools import lru_cache
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from matplotlib import font_manager  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402
from PIL import Image as PILImage  # noqa: E402
from PIL import ImageDraw, ImageFont  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN  # noqa: E402
from pptx.oxml import parse_xml  # noqa: E402
from pptx.oxml.ns import nsdecls, qn  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

from ..config import settings  # noqa: E402

_TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "templates", "reporte_template.pptx")
_FONTS_DIR = os.path.join(os.path.dirname(__file__), "..", "templates", "fonts")

# Paleta institucional (idéntica a la de la plantilla).
NAVY = "#1F3864"
AZUL = "#2E75B6"
CELESTE = "#9DC3E6"
ROJO = "#C00000"

_MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]

FGM = "Franklin Gothic Medium"
TREBUCHET = "Trebuchet MS"
CALIBRI = "Calibri"

# ---------------------------------------------------------------------------
#  Métrica tipográfica real (PIL sobre los .ttf embebidos). Es la MISMA métrica
#  que usará LibreOffice en el contenedor (mismos archivos de fuente), así que
#  lo que aquí cabe, cabe en el PDF final.
# ---------------------------------------------------------------------------
_FONT_FILES = {
    (FGM, False): "framd.ttf",
    (FGM, True): "framd.ttf",           # FGM no tiene bold real (se emboldece sintético)
    (TREBUCHET, False): "trebuc.ttf",
    (TREBUCHET, True): "trebucbd.ttf",
    (CALIBRI, False): "calibri.ttf",
    (CALIBRI, True): "calibrib.ttf",
    ("Times New Roman", False): "times.ttf",
    ("Times New Roman", True): "timesbd.ttf",
}


@lru_cache(maxsize=512)
def _pil_font(family: str, bold: bool, px: int):
    fname = _FONT_FILES.get((family, bold)) or _FONT_FILES.get((family, False))
    if not fname:
        fname = "calibri.ttf"
    try:
        return ImageFont.truetype(os.path.join(_FONTS_DIR, fname), px)
    except Exception:
        return None


def _text_w(text: str, family: str, pt: float, bold: bool = False) -> float:
    """Ancho del texto en PULGADAS a tamaño `pt` (métrica real del .ttf; fallback heurístico)."""
    if not text:
        return 0.0
    font = _pil_font(family, bold, 64)
    if font is None:
        return len(text) * 0.55 * pt / 72
    w64 = font.getlength(text)  # px con tamaño 64 (1px == 1pt a 72dpi)
    w_pt = w64 * pt / 64
    if bold and _FONT_FILES.get((family, bold)) == _FONT_FILES.get((family, False)):
        w_pt *= 1.03  # bold sintético ensancha un poco
    return w_pt / 72


def _fit_pt(text: str, family: str, max_w: float, pt: float, min_pt: float = 6.0,
            bold: bool = False) -> float:
    """Mayor tamaño <= pt con el que `text` cabe en `max_w` pulgadas (sin partir palabras)."""
    while pt > min_pt and _text_w(text, family, pt, bold) > max_w:
        pt -= 0.5
    return pt


def _wrap_lines(text: str, family: str, pt: float, max_w: float, bold: bool = False) -> int:
    """Nº de líneas al envolver `text` por palabras en `max_w` pulgadas (aprox. LibreOffice)."""
    if not text.strip() or max_w <= 0.05:
        return 1
    lines, cur = 1, 0.0
    space = _text_w(" ", family, pt, bold)
    for word in text.split():
        w = _text_w(word, family, pt, bold)
        if w > max_w:  # palabra más ancha que la celda: se parte en trozos
            if cur > 0:
                lines += 1
            lines += int(w // max_w)
            cur = w % max_w
            continue
        if cur > 0 and cur + space + w > max_w:
            lines += 1
            cur = w
        else:
            cur = w if cur == 0 else cur + space + w
    return lines


def _registrar_fuentes_mpl() -> None:
    try:
        for f in set(_FONT_FILES.values()):
            ruta = os.path.join(_FONTS_DIR, f)
            if os.path.exists(ruta):
                font_manager.fontManager.addfont(ruta)
        plt.rcParams["font.family"] = [CALIBRI, "DejaVu Sans"]
    except Exception:
        pass


_registrar_fuentes_mpl()


# ---------------------------------------------------------------------------
#  Helpers de datos / formato
# ---------------------------------------------------------------------------
def _get(obj: Any, path: str) -> Any:
    """Resuelve rutas tipo 'a.b[0].c' dentro del contenido."""
    cur = obj
    for part in path.split("."):
        m = re.match(r"^([^\[]+)(?:\[(\d+)\])?$", part)
        if not m:
            return None
        key, idx = m.group(1), m.group(2)
        cur = cur.get(key) if isinstance(cur, dict) else None
        if idx is not None:
            cur = cur[int(idx)] if isinstance(cur, list) and len(cur) > int(idx) else None
        if cur is None:
            return None
    return cur


def _clean(v: Any) -> str:
    """Valor legible o '' si es placeholder ('-.-', '- . -', vacío, guiones)."""
    if v is None:
        return ""
    s = str(v).strip()
    if not s or set(s) <= {"-", " ", ".", "–", "—"}:
        return ""
    return s


def _num(v: Any) -> float | None:
    """'2.305'/'5.911'/'4.565 Mts'/50 -> float (separador de miles '.'); si no, None."""
    if v is None:
        return None
    s = str(v).strip().replace(".", "").replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return None
    try:
        return float(m.group())
    except ValueError:
        return None


def _miles(v: float | None) -> str:
    """43720 -> '43.720' (separador de miles institucional)."""
    if v is None:
        return ""
    return f"{int(round(v)):,}".replace(",", ".")


def _fecha_es(f: datetime) -> str:
    """04JUL2026 (día + mes ES abreviado + año), igual que el encabezado de la plantilla."""
    return f"{f.day:02d}{_MESES[f.month - 1]}{f.year}"


def _by_name(slide) -> dict[str, Any]:
    return {sh.name: sh for sh in slide.shapes}


def _v(contenido: dict, path: str) -> str:
    return _clean(_get(contenido, path))


def _items(contenido: dict, path: str) -> list[dict]:
    return [it for it in (_get(contenido, path) or []) if isinstance(it, dict)]


def _dget(contenido: dict, path: str) -> dict:
    """Como _get, pero garantiza dict (el LLM a veces devuelve str/list donde se espera un objeto)."""
    v = _get(contenido, path)
    return v if isinstance(v, dict) else {}


# ---------------------------------------------------------------------------
#  Ajuste de tipografía / geometría
# ---------------------------------------------------------------------------
def _fit(tf, pt: float | None = None, wrap: bool = True) -> None:
    """Normaliza un text_frame: word_wrap, sin autoajuste y (si `pt`) tamaño en todos los runs."""
    tf.word_wrap = wrap
    try:
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    if pt is not None:
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(pt)


def _shape_fit(shape, pt: float | None = None, wrap: bool = True) -> None:
    if shape is not None and shape.has_text_frame:
        _fit(shape.text_frame, pt, wrap)


def _para_fit(tf, idx: int, pt: float) -> None:
    """Fija el tamaño de un párrafo concreto (para KPIs que mezclan rótulo pequeño + cifra grande)."""
    paras = tf.paragraphs
    if -len(paras) <= idx < len(paras):
        for r in paras[idx].runs:
            r.font.size = Pt(pt)


def _move(shape, left=None, top=None, width=None, height=None) -> None:
    if left is not None:
        shape.left = left
    if top is not None:
        shape.top = top
    if width is not None:
        shape.width = width
    if height is not None:
        shape.height = height


def _to_front(shape) -> None:
    """Trae el shape al frente (después de insertar un PNG encima de su zona)."""
    sp = shape._element
    spTree = sp.getparent()
    spTree.remove(sp)
    spTree.append(sp)


def _center_title(shape, pt: float = 18.0, family: str = FGM) -> None:
    """Títulos verdes de la portada: quedan a `pt` (Franklin Gothic Medium 18) en UNA línea.

    La caja se re-centra al ancho real del texto para que LibreOffice no lo envuelva ni
    lo desplace: mismo centro, ancho = ancho medido + margen."""
    if shape is None or not shape.has_text_frame:
        return
    tf = shape.text_frame
    _fit(tf, pt=pt, wrap=False)
    texto = " ".join(tf.text.split())
    w = Inches(_text_w(texto, family, pt) + 0.25)
    cx = shape.left + shape.width // 2
    shape.left = int(cx - w // 2)
    shape.width = int(w)


# ---------------------------------------------------------------------------
#  Edición de texto / tablas conservando formato
# ---------------------------------------------------------------------------
def _set_para(tf, idx: int, text: str) -> None:
    """Edita el párrafo `idx`: fija runs[0].text y elimina runs sobrantes (conserva formato)."""
    paras = tf.paragraphs
    if idx < 0:
        idx += len(paras)
    if not (0 <= idx < len(paras)):
        return
    p = paras[idx]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run().text = text


def _set_text(tf, text: str) -> None:
    """Colapsa el text_frame a un solo párrafo/run con `text`.

    Conserva el formato del primer párrafo QUE TENGA runs: algunas celdas de la plantilla
    parten con un párrafo vacío (sin run) y usarlo heredaría la fuente default (~18pt
    Calibri) en vez de la de la celda."""
    paras = list(tf.paragraphs)
    keep = None
    for p in paras:
        if p.runs:
            keep = p
            break
    if keep is None:
        keep = paras[0]
    for p in paras:
        if p._p is not keep._p:
            p._p.getparent().remove(p._p)
    _set_para(tf, 0, text)


def _append_line(tf, text: str) -> None:
    """Añade una línea nueva clonando el formato del último párrafo (para captions KPI)."""
    if not text:
        return
    last = tf.paragraphs[-1]._p
    new = deepcopy(last)
    last.getparent().append(new)
    _set_para(tf, len(tf.paragraphs) - 1, text)


def _set_cell(cell, text: str, size: float | None = None, bold: bool | None = None,
              align=None, font: str | None = None, underline: bool | None = None) -> None:
    _set_text(cell.text_frame, text)
    if size is not None:
        _fit(cell.text_frame, size, wrap=True)
    else:
        _fit(cell.text_frame, None, wrap=True)
    for p in cell.text_frame.paragraphs:
        if align is not None:
            p.alignment = align
        for r in p.runs:
            if bold is not None:
                r.font.bold = bold
            if font is not None:
                r.font.name = font
            if underline is not None:
                r.font.underline = underline


def _cell_value_line(cell, value: str, size: float | None = None, bold: bool = True) -> None:
    """Celdas 'ETIQUETA' + valor DEBAJO (2ª línea) conservando la etiqueta de la plantilla."""
    tf = cell.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    if not value:
        return
    _append_line(tf, value)
    p = tf.paragraphs[-1]
    for r in p.runs:
        if size is not None:
            r.font.size = Pt(size)
        r.font.bold = bold
    _fit(tf, None, wrap=True)


def _clone_rows(tbl, at_index: int, count: int) -> None:
    """Clona la fila <a:tr> `at_index` `count` veces, insertándola a continuación."""
    trs = tbl._tbl.findall(qn("a:tr"))
    src = trs[at_index]
    for _ in range(count):
        src.addnext(deepcopy(src))


def _delete_rows(tbl, start: int, count: int) -> None:
    """Elimina `count` filas <a:tr> desde `start` (filas de ejemplo que quedaron sin dato)."""
    trs = tbl._tbl.findall(qn("a:tr"))
    for tr in trs[start:start + count]:
        tr.getparent().remove(tr)


_LN_EDGES = ("lnL", "lnR", "lnT", "lnB")


def _set_cell_border(cell, edge: str, color: str | None) -> None:
    """Borde EXPLÍCITO de una celda (`edge` in lnL/lnR/lnT/lnB): sólido `color` o noFill
    si es None. Explícito porque donde el borde no está definido LibreOffice puede aplicar
    el default del estilo de tabla (línea fantasma que PowerPoint no muestra)."""
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn(f"a:{edge}"))
    if old is not None:
        tcPr.remove(old)
    if color is None:
        ln = parse_xml(f'<a:{edge} {nsdecls("a")} w="12700"><a:noFill/></a:{edge}>')
    else:
        ln = parse_xml(
            f'<a:{edge} {nsdecls("a")} w="12700" cap="flat">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:prstDash val="solid"/></a:{edge}>'
        )
    # Los ln* van al inicio de tcPr y en orden lnL,lnR,lnT,lnB (schema DrawingML).
    prev = None
    for name in _LN_EDGES[:_LN_EDGES.index(edge)]:
        e = tcPr.find(qn(f"a:{name}"))
        if e is not None:
            prev = e
    if prev is not None:
        prev.addnext(ln)
    else:
        tcPr.insert(0, ln)


def _unmerge_col(shape, col: int, rows: list[int]) -> None:
    """Deshace el merge vertical de una columna (p.ej. COMANDO MATRIZ): cada fila su celda."""
    trs = shape.table._tbl.findall(qn("a:tr"))
    for r in rows:
        if r >= len(trs):
            continue
        tcs = trs[r].findall(qn("a:tc"))
        if col < len(tcs):
            tc = tcs[col]
            for attr in ("rowSpan", "vMerge"):
                if tc.get(attr):
                    del tc.attrib[attr]


def _set_col_widths(shape, widths: dict[int, int]) -> None:
    """Reasigna anchos de columna (EMU). La suma debe conservar el ancho total de la tabla."""
    for i, col in enumerate(shape.table.columns):
        if i in widths:
            col.width = widths[i]


def _fill_table_list(shape, items: list[dict], start_row: int, cols: dict[int, str],
                     footer: int = 0, can_clone: bool = True, size: float | None = None,
                     max_rows: int | None = None, font: str | None = None,
                     underline: bool | None = None,
                     aligns: dict[int, Any] | None = None) -> None:
    """Rellena una tabla repetitiva desde `start_row`. ELIMINA las filas de ejemplo sobrantes
    (sin dato no queda fila: nada de filas vacías en el reporte).

    `cols`: {índice_columna: campo_json}. `footer`: nº de filas finales intocables (p.ej. TOTAL).
    `size`: tamaño de fuente (pt) aplicado a las celdas de datos; imprescindible porque las
    celdas de ejemplo vienen vacías (sin run) y el run nuevo heredaría el default gigante (~18pt).
    `aligns`: {índice_columna: PP_ALIGN} para uniformar la alineación (las filas nuevas no
    heredan la de la fila de ejemplo del template).
    """
    tbl = shape.table
    items = [it for it in (items or []) if isinstance(it, dict)]
    if max_rows is not None:
        items = items[:max_rows]
    n = len(items)
    end = len(tbl.rows) - footer          # fin (exclusivo) de la zona de datos
    avail = end - start_row
    if n > avail and can_clone and avail >= 1:
        _clone_rows(tbl, at_index=end - 1, count=n - avail)
        end += n - avail
    for i in range(min(n, end - start_row)):
        for c, field in cols.items():
            _set_cell(tbl.cell(start_row + i, c), _clean(items[i].get(field)), size,
                      font=font, underline=underline, align=(aligns or {}).get(c))
    # Filas de ejemplo que quedaron sin dato: se eliminan (tampoco mostrar cifras horneadas).
    surplus = end - (start_row + n)
    if surplus > 0:
        _delete_rows(tbl, start_row + n, surplus)


def _fill_cells(shape, cellmap: dict[tuple[int, int], str], size: float | None = None) -> None:
    tbl = shape.table
    for (r, c), val in cellmap.items():
        _set_cell(tbl.cell(r, c), val, size)


def _fill_labelvalue(shape, cellmap: dict[tuple[int, int], str], size: float | None = None) -> None:
    """Celdas 'ETIQUETA: valor': conserva la etiqueta del template y le anexa el valor.

    Las celdas de rótulo de la plantilla traen algn="r" con un marR enorme (el rótulo
    terminaba donde empezaba el valor manuscrito); con el valor anexado ese margen hace
    envolver el texto -> se elimina y se centra el par etiqueta+valor sobre la línea."""
    tbl = shape.table
    for (r, c), val in cellmap.items():
        cell = tbl.cell(r, c)
        label = " ".join(cell.text.split())
        _set_cell(cell, f"{label} {val}".strip() if val else label, size)
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            pPr = p._p.find(qn("a:pPr"))
            if pPr is not None and pPr.get("marR"):
                del pPr.attrib["marR"]


def _clear_cells(shape, coords: list[tuple[int, int]]) -> None:
    tbl = shape.table
    for r, c in coords:
        _set_cell(tbl.cell(r, c), "")


def _cell_fit(shape, sizes: dict[tuple[int, int], float], wrap: bool = False) -> None:
    """Reajusta el tamaño de celdas concretas (p.ej. títulos de banda horneados en la tabla)."""
    tbl = shape.table
    for (r, c), pt in sizes.items():
        _fit(tbl.cell(r, c).text_frame, pt, wrap)


# ---------------------------------------------------------------------------
#  Estimación de altura de tablas (para el reflow: LibreOffice AGRANDA las filas
#  cuyo texto no cabe, empujando la tabla hacia abajo; hay que mover lo de abajo).
# ---------------------------------------------------------------------------
def _est_table_h(shape) -> int:
    """Altura estimada (EMU) de la tabla tras el relleno, con la métrica real de fuentes."""
    tbl = shape.table
    col_w = [c.width for c in tbl.columns]
    trs = tbl._tbl.findall(qn("a:tr"))
    total = 0
    for ri, tr in enumerate(trs):
        base = int(tr.get("h") or 0)
        content = 0
        tcs = tr.findall(qn("a:tc"))
        for ci, tc in enumerate(tcs):
            if ci >= len(col_w) or tc.get("hMerge") or tc.get("vMerge"):
                continue
            span = int(tc.get("gridSpan") or 1)
            w_in = sum(col_w[ci:ci + span]) / 914400.0 - 0.17  # insets izq+der
            try:
                cell = tbl.cell(ri, ci)
            except Exception:
                continue
            h_in = 0.0
            for p in cell.text_frame.paragraphs:
                texto = "".join(r.text for r in p.runs)
                if not texto.strip():
                    continue
                r0 = p.runs[0]
                pt = r0.font.size.pt if r0.font.size else 18.0
                fam = r0.font.name or CALIBRI
                bold = bool(r0.font.bold)
                lines = _wrap_lines(texto, fam, pt, max(w_in, 0.25), bold)
                h_in += lines * pt * 1.26 / 72
            if h_in:
                h_in += 0.10  # insets sup+inf
            content = max(content, int(h_in * 914400))
        total += max(base, content)
    return total


def _bottom(shape) -> int:
    return int(shape.top + _est_table_h(shape)) if shape.has_table else int(shape.top + shape.height)


# ---------------------------------------------------------------------------
#  Gráficos (matplotlib -> PNG)
# ---------------------------------------------------------------------------
def _fig_png(fig, transparent: bool = True, facecolor=None) -> bytes:
    buf = io.BytesIO()
    kw: dict[str, Any] = {"format": "png", "dpi": 200, "bbox_inches": "tight"}
    if facecolor is not None:
        kw["facecolor"] = facecolor
        kw["transparent"] = False
    else:
        kw["transparent"] = transparent
    fig.savefig(buf, **kw)
    plt.close(fig)
    return buf.getvalue()


def _planta_variable(pf: dict) -> tuple[float, float, float]:
    planta = sum(x for k in ("OF", "SOF", "ECP", "ESCMIL", "ESCSOF", "ESCSERV", "personal_civil")
                 if (x := _num(pf.get(k))))
    variable = sum(x for k in ("SLTP", "SLC") if (x := _num(pf.get(k))))
    total = _num(pf.get("total")) or (planta + variable)
    return planta, variable, total


def _chart_donut(pf: dict) -> bytes | None:
    """Dona Pers. Variable (SLTP+SLC) vs Planta (resto). Centro = total. (Página 2)."""
    planta, variable, total = _planta_variable(pf)
    if planta + variable <= 0:
        return None
    fig, ax = plt.subplots(figsize=(3.1, 3.1))
    ax.pie(
        [planta, variable],
        labels=[f"Planta\n{_miles(planta)}", f"Variable\n{_miles(variable)}"],
        colors=[NAVY, CELESTE],
        startangle=90,
        wedgeprops={"width": 0.42, "edgecolor": "white"},
        textprops={"fontsize": 9, "color": "#333"},
    )
    ax.text(0, 0.07, _miles(total), ha="center", va="center", fontsize=15,
            fontweight="bold", color=NAVY, fontfamily=TREBUCHET)
    ax.text(0, -0.20, "TOTAL", ha="center", va="center", fontsize=8, color="#555")
    ax.set(aspect="equal")
    return _fig_png(fig)


def _chart_mini_donut(pf: dict) -> bytes | None:
    """Mini dona de la PORTADA (dentro de PARTE DE FUERZA INSTITUCIONAL), total al centro."""
    planta, variable, total = _planta_variable(pf)
    if planta + variable <= 0:
        return None
    fig, ax = plt.subplots(figsize=(1.6, 1.6))
    ax.pie([planta, variable], colors=[NAVY, CELESTE], startangle=90,
           wedgeprops={"width": 0.38, "edgecolor": "white"})
    ax.text(0, 0, _miles(total), ha="center", va="center", fontsize=10.5,
            fontweight="bold", color=NAVY, fontfamily=TREBUCHET)
    ax.set(aspect="equal")
    return _fig_png(fig)


def _chart_macrozonas(macrozonas: list[dict]) -> bytes | None:
    """Barras horizontales de fuerza por macrozona (página 2, sección 2)."""
    pares = [(m.get("zona", ""), _num(m.get("fuerza"))) for m in (macrozonas or []) if isinstance(m, dict)]
    pares = [(z, v) for z, v in pares if z and v]
    if not pares:
        return None
    fig, ax = plt.subplots(figsize=(5.4, 2.5))
    zs = [z for z, _ in pares]
    vs = [v for _, v in pares]
    barras = ax.barh(zs, vs, color=NAVY, height=0.62)
    ax.invert_yaxis()
    ax.set_xlim(0, max(vs) * 1.2)
    ax.bar_label(barras, labels=[_miles(v) for v in vs], padding=4, fontsize=9, color="#333")
    ax.tick_params(labelsize=9)
    ax.tick_params(axis="x", labelsize=8, colors="#555")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return _fig_png(fig)


def _chart_catastrofe(cat: dict) -> bytes | None:
    """Barras FUERZA/FORMAN/FALTAN de la fuerza de catástrofe (página 2, sección 3)."""
    filas = [(lbl, _num(cat.get(k))) for lbl, k in
             (("FUERZA", "fuerza"), ("FORMAN", "forman"), ("FALTAN", "faltan"))]
    filas = [(l, v) for l, v in filas if v is not None]
    if not filas:
        return None
    colores = {"FUERZA": NAVY, "FORMAN": AZUL, "FALTAN": ROJO}
    fig, ax = plt.subplots(figsize=(5.4, 1.9))
    ls = [l for l, _ in filas]
    vs = [v for _, v in filas]
    barras = ax.barh(ls, vs, color=[colores[l] for l in ls], height=0.6)
    ax.invert_yaxis()
    ax.set_xlim(0, max(vs) * 1.22)
    ax.bar_label(barras, labels=[_miles(v) for v in vs], padding=4, fontsize=9, color="#333")
    ax.tick_params(labelsize=9)
    ax.tick_params(axis="x", labelsize=8, colors="#555")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return _fig_png(fig)


def _chart_uac(items: list[dict]) -> bytes | None:
    """Traslados por UAC (página 5): tarjeta azul con barras celestes y rótulos blancos."""
    pares = [(_clean(it.get("uac")), _num(it.get("unidades"))) for it in items or []
             if isinstance(it, dict)]
    pares = [(u, v) for u, v in pares if u and v is not None]
    if not pares:
        return None
    fig = plt.figure(figsize=(6.15, 3.13))
    fig.patch.set_facecolor(NAVY)
    # Zona superior libre para el título "Traslados por UAC" (textbox de la plantilla).
    ax = fig.add_axes([0.035, 0.40, 0.93, 0.34])
    ax.set_facecolor(NAVY)
    x = np.arange(len(pares))
    vs = [v for _, v in pares]
    barras = ax.bar(x, vs, width=0.45, color=CELESTE)
    ax.bar_label(barras, labels=[_miles(v) for v in vs], padding=2, fontsize=9.5, color="white")
    ax.set_xticks(x)
    ax.set_xticklabels([u for u, _ in pares], rotation=38, ha="right", fontsize=8.5, color="white")
    ax.set_ylim(0, max(vs) * 1.3)
    ax.set_yticks([])
    for s in ax.spines.values():
        s.set_visible(False)
    ax.spines["bottom"].set_visible(True)
    ax.spines["bottom"].set_color("#B7C4DB")
    ax.tick_params(axis="x", length=0, pad=1)
    fig.legend(handles=[Patch(facecolor=CELESTE)], labels=["UNIDADES"], loc="lower center",
               frameon=False, fontsize=8.5, labelcolor="white", handlelength=0.9,
               handleheight=0.9, bbox_to_anchor=(0.5, -0.015))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, facecolor=NAVY)  # sin bbox tight: conserva la tarjeta completa
    plt.close(fig)
    return buf.getvalue()


def _chart_operacionalidad(items: list[dict], resumen: dict) -> bytes | None:
    """Barras agrupadas Total vs NOP (no operativos) por unidad (página 6)."""
    filas = []
    for it in items or []:
        if not isinstance(it, dict):
            continue
        u = _clean(it.get("unidad"))
        t, n = _num(it.get("total")), _num(it.get("nop"))
        if u and (t or n):
            filas.append((u, t or 0, n or 0))
    if not filas:
        return None
    fig, ax = plt.subplots(figsize=(7.4, 5.2))
    x = np.arange(len(filas))
    w = 0.38
    b1 = ax.bar(x - w / 2, [f[1] for f in filas], w, label="Total", color=NAVY)
    b2 = ax.bar(x + w / 2, [f[2] for f in filas], w, label="NOP", color=ROJO)
    ax.bar_label(b1, fmt=lambda v: _miles(v), fontsize=7.5, padding=2, color="#333")
    ax.bar_label(b2, fmt=lambda v: _miles(v), fontsize=7.5, padding=2, color=ROJO)
    ax.set_xticks(x)
    ax.set_xticklabels([f[0] for f in filas], rotation=30, ha="right", fontsize=8.5)
    ax.legend(fontsize=9.5)
    ax.tick_params(axis="y", labelsize=8)
    titulo = "Operacionalidad logística por unidad"
    tot, nop = _clean((resumen or {}).get("total")), _clean((resumen or {}).get("nop"))
    if tot or nop:
        titulo += f"   (Total {tot or '-'} / NOP {nop or '-'})"
    ax.set_title(titulo, fontsize=11.5, color=NAVY, fontweight="bold")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return _fig_png(fig)


def _reemplazar(slide, names: list[str], png: bytes | None, rect=None) -> None:
    """Elimina los shapes `names` (formas nativas del gráfico) e inserta `png` en su
    bounding box (o en `rect`). Si `png` es None, solo elimina: el gráfico se OMITE por
    completo (no se deja la forma de ejemplo del template con su escala horneada)."""
    objetivo = set(names)
    shapes = [sh for sh in slide.shapes if sh.name in objetivo]
    if not shapes and rect is None:
        return
    if rect is None and shapes:
        left = min(s.left for s in shapes)
        top = min(s.top for s in shapes)
        width = max(s.left + s.width for s in shapes) - left
        height = max(s.top + s.height for s in shapes) - top
        rect = (left, top, width, height)
    for s in shapes:
        s._element.getparent().remove(s._element)
    if png is not None and rect is not None:
        slide.shapes.add_picture(io.BytesIO(png), *rect)


# Furniture del gráfico de barras por macrozona (grupo + ejes/labels) y del gráfico de
# catástrofe (ejes 0..800 abajo a la derecha) en la página 2 (slide idx1).
_MACRO_NAMES = ["object 28"] + [f"object {i}" for i in range(41, 56)]
_CATA_NAMES = [f"object {i}" for i in range(56, 67)]
# Gráfico "Traslados por UAC" horneado en la página 5 (slide idx4): tarjeta con barras
# dibujadas + cajas de valores + iconos de ejes + leyenda. Se regenera completo.
_UAC_NAMES = ["object 2", "object 7", "object 8", "object 9", "object 10", "object 11",
              "object 12", "object 21"]


# ---------------------------------------------------------------------------
#  Abreviaciones institucionales (celdas angostas de la portada)
# ---------------------------------------------------------------------------
_JAF_CORTAS = {
    "arica y parinacota": "ARICA PARIN",
    "tarapaca": "TARAP",
    "tarapacá": "TARAP",
    "antofagasta": "ANTOF",
}


def _jaf_corta(nombre: str) -> str:
    n = _clean(nombre)
    base = re.sub(r"^\s*JAF\s+", "", n, flags=re.I).strip()
    return _JAF_CORTAS.get(base.lower(), base.upper()[:11])


# ---------------------------------------------------------------------------
#  PORTADA (slide idx0)
# ---------------------------------------------------------------------------
def _s0_portada(slide, c: dict) -> None:
    s0 = _by_name(slide)
    pf = "personal.parte_fuerza_institucional"

    # --- FUERZA SLC ---
    _fill_cells(s0["object 34"], {
        (1, 1): _v(c, "personal.slc.clase_2006"),
        (2, 1): _v(c, "personal.slc.clase_2007"),
        (3, 1): _v(c, "personal.slc.total"),
    }, size=9)

    # --- PARTE DE FUERZA INSTITUCIONAL: el total va DENTRO de la mini-dona (PNG).
    #     Las filas 4 y 5 son slivers vacíos (solo traían el '25,672' de ejemplo) -> fuera.
    #     La línea que separa ECP/SLC de SLTP/PERSONAL CIVIL se conserva SOLO bajo las
    #     columnas de rótulos; en la columna 0 se anula todo borde interno para que
    #     ninguna línea atraviese la dona (LibreOffice dibujaba la de los slivers encima).
    t39 = s0["object 39"].table
    _delete_rows(t39, 4, 2)
    _clear_cells(s0["object 39"], [(3, 0)])
    _set_cell(t39.cell(4, 2), "PERSONAL CIVIL", size=8)
    _set_cell_border(t39.cell(3, 1), "lnB", "000000")
    _set_cell_border(t39.cell(3, 2), "lnB", "000000")
    _set_cell_border(t39.cell(2, 0), "lnB", None)
    _set_cell_border(t39.cell(3, 0), "lnT", None)
    _set_cell_border(t39.cell(3, 0), "lnB", None)
    _set_cell_border(t39.cell(4, 0), "lnT", None)

    # --- DESPLEGADOS EN MZs (suma de macrozonas) / OTRAS MOOTW (catástrofe) /
    #     OTROS DESPLIEGUES (Antártica y OPAZ) ---
    t33 = s0["object 33"].table
    mz = _items(c, "personal.macrozonas")
    sumas = {k: sum(_num(m.get(k)) or 0 for m in mz) for k in ("fuerza", "forman", "faltan")}
    if any(sumas.values()):
        for col, k in enumerate(("fuerza", "forman", "faltan")):
            _cell_value_line(t33.cell(1, col), _miles(sumas[k]), size=8.5)
    cat = _dget(c, "personal.catastrofe")
    if any(_clean(cat.get(k)) for k in ("fuerza", "forman", "faltan")):
        for col, k in enumerate(("fuerza", "forman", "faltan")):
            _cell_value_line(t33.cell(3, col), _clean(cat.get(k)), size=8.5)
    antart = _items(c, "operaciones.ops_antartica")
    dot = sum(_num(a.get("dotacion_estival")) or 0 for a in antart)
    _set_cell(t33.cell(5, 2), _miles(dot) if dot else "", size=9, bold=True,
              align=PP_ALIGN.CENTER)
    extr = _items(c, "operaciones.ops_extranjero")
    opaz = sum(_num(e.get("efectivos")) or 0 for e in extr)
    _set_cell(t33.cell(6, 2), _miles(opaz) if opaz else "", size=9, bold=True,
              align=PP_ALIGN.CENTER)
    # Centrado vertical: los valores quedaban pegados al borde superior de la celda.
    t33.cell(5, 2).vertical_anchor = MSO_ANCHOR.MIDDLE
    t33.cell(6, 2).vertical_anchor = MSO_ANCHOR.MIDDLE

    # --- Incidentes relevantes (MOOTW): UNIDAD | TIPO ---
    _fill_table_list(s0["object 40"], _items(c, "inteligencia.incidentes_mootw"),
                     start_row=1, cols={0: "unidad", 1: "tipo"}, can_clone=True,
                     size=10, max_rows=3)

    # --- Unidades de emergencia: UFEC / BRIFE / PARME ---
    ue = _dget(c, "personal.unidades_emergencia")
    _fill_cells(s0["object 3"], {
        (1, 0): _clean(ue.get("ufec")), (1, 1): _clean(ue.get("brife")),
        (1, 2): _clean(ue.get("parme")),
    }, size=10)

    # --- Relevos: N.º y periodo por macrozona ---
    t7 = s0["object 7"].table
    rn = _items(c, "operaciones.relevos_norte")
    rs = _items(c, "operaciones.relevos_sur")
    _set_cell(t7.cell(1, 0), f"N.º {_clean(rn[0].get('n'))}" if rn and _clean(rn[0].get("n")) else "",
              size=9, bold=True)
    _set_cell(t7.cell(1, 1), (_clean(rn[0].get("periodo")) or _clean(rn[0].get("fecha"))) if rn else "",
              size=9, bold=True)
    _set_cell(t7.cell(3, 0), f"N.º {_clean(rs[0].get('n'))}" if rs and _clean(rs[0].get("n")) else "",
              size=9, bold=True)
    _set_cell(t7.cell(3, 1), (_clean(rs[0].get("periodo")) or _clean(rs[0].get("fecha"))) if rs else "",
              size=9, bold=True)

    # --- Ops Guardián Soberano (JAF abreviada para la columna angosta) ---
    guard = [{"jaf": _jaf_corta(g.get("jaf", "")), "fuerza": g.get("fuerza"),
              "planif": g.get("planif"), "ejecut": g.get("ejecut")}
             for g in _items(c, "operaciones.guardian_soberano")]
    _fill_table_list(s0["object 10"], guard, start_row=2,
                     cols={0: "jaf", 1: "fuerza", 2: "planif", 3: "ejecut"}, size=7)

    # --- Ops. en el extranjero: UNIDAD (operación + país) | FUERZA + TOTAL ---
    ext_rows = [{"unidad": f"{_clean(e.get('operacion'))} ({_clean(e.get('ubicacion'))})"
                 if _clean(e.get("ubicacion")) else _clean(e.get("operacion")),
                 "fuerza": e.get("efectivos")} for e in extr]
    t8 = s0["object 8"]
    _fill_table_list(t8, ext_rows, start_row=1, cols={0: "unidad", 1: "fuerza"},
                     footer=1, can_clone=False, size=8)
    _set_cell(t8.table.cell(len(t8.table.rows) - 1, 1), _miles(opaz) if opaz else "",
              size=9, bold=True, align=PP_ALIGN.CENTER)

    # --- Unidades en terreno (5 columnas) + TOTAL ---
    terreno = _items(c, "operaciones.unidades_terreno")
    t5 = s0["object 5"]
    _set_col_widths(t5, {0: Inches(0.85), 1: Inches(0.90), 2: Inches(0.93),
                         3: Inches(0.89), 4: Inches(0.54)})
    _fill_table_list(t5, terreno, start_row=1,
                     cols={0: "comando_matriz", 1: "ur", 2: "fechas", 3: "ubicacion", 4: "fuerza"},
                     footer=1, can_clone=False, size=7.5)
    tot_f = sum(_num(u.get("fuerza")) or 0 for u in terreno)
    _set_cell(t5.table.cell(len(t5.table.rows) - 1, 4), _miles(tot_f) if tot_f else "",
              size=8, bold=True, align=PP_ALIGN.CENTER)

    # --- Títulos verdes: Franklin Gothic Medium 18, UNA línea, caja re-centrada ---
    for t in ("object 2", "object 4", "object 6", "object 9", "object 11", "object 12"):
        _center_title(s0.get(t), pt=18)

    # --- Badges del mapa: MZN / MZS / ANTÁRTICA + nº de unidades desplegadas ---
    mapa = _dget(c, "operaciones.despliegue_mapa")
    for name, label, lbl_pt, key in (("object 24", "MZN", 12, "mzn"),
                                     ("object 32", "MZS", 12, "mzs"),
                                     ("object 25", "ANTÁRTICA", 10.5, "antartica")):
        sh = s0.get(name)
        if sh is None:
            continue
        tf = sh.text_frame
        _set_para(tf, 0, label)
        _set_para(tf, 1, _clean(mapa.get(key)))
        _fit(tf, None, wrap=False)
        w = Inches(max(_text_w(label, FGM, lbl_pt) + 0.10, sh.width / 914400))
        cx = sh.left + sh.width // 2
        sh.left = int(cx - w // 2)
        sh.width = int(w)

    # --- Reflow vertical: LibreOffice AGRANDA las filas que no caben, así que cada
    #     bloque baja hasta quedar bajo el anterior (título verde + su tabla juntos).
    #     Cadena: tablas superiores -> título Incidentes + tabla -> columna central
    #     (U. emergencia -> U. terreno) y columna derecha (Relevos -> Guardián -> Ops ext.).
    def _bloque(titulo: str, tabla: str, top_nominal: float, tope: int, gap: float):
        top = max(Inches(top_nominal), tope + Inches(gap))
        offset = int(Inches(0.35))
        sh_t, sh_b = s0.get(titulo), s0[tabla]
        if sh_t is not None:
            offset = sh_b.top - sh_t.top
            _move(sh_t, top=int(top))
        _move(sh_b, top=int(top) + offset)
        return _bottom(sh_b)

    b_sup = max(_bottom(s0["object 39"]), _bottom(s0["object 33"]),
                _bottom(s0["object 34"]))
    b40 = _bloque("object 12", "object 40", 3.11, b_sup, 0.10)   # Incidentes relevantes
    b3 = _bloque("object 2", "object 3", 4.08, b40, 0.10)        # Unidades de emergencia
    _bloque("object 4", "object 5", 5.58, b3, 0.12)              # Unidades en terreno
    b7 = _bloque("object 6", "object 7", 4.09, b40, 0.10)        # Relevos
    b10 = _bloque("object 11", "object 10", 5.58, b7, 0.12)      # Ops Guardián Soberano
    _bloque("object 9", "object 8", 7.77, b10, 0.12)             # Ops. en el extranjero


# ---------------------------------------------------------------------------
#  PERSONAL (slide idx1)
# ---------------------------------------------------------------------------
def _s1_personal(slide, c: dict, fecha: datetime) -> None:
    s1 = _by_name(slide)
    pf = "personal.parte_fuerza_institucional"

    # La columna izquierda (OF:/SOF:/ECP:/SLTP:) es angosta y el valor saltaba de línea:
    # se ensancha a costa de la derecha (que tiene holgura de sobra).
    _set_col_widths(s1["object 2"], {0: Inches(1.55), 1: Inches(2.68)})
    _fill_labelvalue(s1["object 2"], {
        (0, 0): _v(c, f"{pf}.OF"), (1, 0): _v(c, f"{pf}.SOF"),
        (2, 0): _v(c, f"{pf}.ECP"), (3, 0): _v(c, f"{pf}.SLTP"),
        (0, 1): _v(c, f"{pf}.ESCMIL"), (1, 1): _v(c, f"{pf}.ESCSOF"),
        (2, 1): _v(c, f"{pf}.ESCSERV"), (3, 1): _v(c, f"{pf}.personal_civil"),
        (4, 1): _v(c, f"{pf}.SLC"),
    }, size=13)

    # FUERZA/FORMAN/FALTAN — PARTE DE FUERZA MACROZONAS (suma de las 4 macrozonas):
    # el template solo trae la fila de rótulos; se agrega la fila de valores con aire
    # (0.44" y centrado vertical: pegada al borde se veía mal). La línea decorativa
    # suelta bajo la tabla ('object 27') sobra y se elimina.
    t26 = s1["object 26"]
    mz = _items(c, "personal.macrozonas")
    sumas = {k: sum(_num(m.get(k)) or 0 for m in mz) for k in ("fuerza", "forman", "faltan")}
    if any(sumas.values()):
        if len(t26.table.rows) == 1:
            _clone_rows(t26.table, at_index=0, count=1)
        t26.table._tbl.findall(qn("a:tr"))[1].set("h", str(int(Inches(0.44))))
        for col, k in enumerate(("fuerza", "forman", "faltan")):
            _set_cell(t26.table.cell(1, col), _miles(sumas[k]) if sumas[k] else "",
                      size=15, bold=True)
            t26.table.cell(1, col).vertical_anchor = MSO_ANCHOR.MIDDLE
    raya = s1.get("object 27")
    if raya is not None:
        raya._element.getparent().remove(raya._element)

    # FUERZA/FORMAN/FALTAN — PARTE DE FUERZA DE CATÁSTROFE (conserva Trebuchet 18 B).
    _fill_cells(s1["object 25"], {
        (1, 0): _v(c, "personal.catastrofe.fuerza"),
        (1, 1): _v(c, "personal.catastrofe.forman"),
        (1, 2): _v(c, "personal.catastrofe.faltan"),
    })

    _set_text(s1["object 20"].text_frame, str(fecha.year))
    _shape_fit(s1["object 20"], pt=15.5, wrap=False)


# ---------------------------------------------------------------------------
#  INTELIGENCIA (slide idx2)
# ---------------------------------------------------------------------------
def _s2_inteligencia(slide, c: dict, fstr: str) -> None:
    s2 = _by_name(slide)

    # Título de la banda 1 (a 15.5pt no cabe en LibreOffice: se fija 14 y caja más ancha).
    tit = s2["object 8"]
    _move(tit, left=Inches(3.30), width=Inches(5.05))
    _shape_fit(tit, pt=14, wrap=False)

    # --- 1. Incidentes institucionales: columna TIPO más ancha (evita partir palabras) ---
    t9 = s2["object 9"]
    _set_col_widths(t9, {0: Inches(1.45), 1: Inches(4.42)})
    _fill_table_list(t9, _items(c, "inteligencia.incidentes_institucionales"),
                     start_row=1, cols={0: "tipo", 1: "descripcion"}, size=9)

    # --- 2. Incidentes MOOTW: columna UNIDAD más ancha ---
    t21 = s2["object 21"]
    _set_col_widths(t21, {2: Inches(1.22), 3: Inches(3.37)})
    _fill_table_list(t21, _items(c, "inteligencia.incidentes_mootw"),
                     start_row=2, cols={0: "tipo", 2: "unidad", 3: "descripcion"}, size=9)

    # --- 3/4/5. Sistema de alerta + meteorología (una sola tabla con bandas) ---
    t10 = s2["object 10"]
    tbl = t10.table
    _cell_fit(t10, {(1, 0): 10, (1, 2): 10, (1, 3): 10, (1, 4): 10}, wrap=True)  # cabeceras 13.5 -> 10
    alertas = _items(c, "inteligencia.sistema_alerta")
    extra = max(0, len(alertas) - 1)
    if extra:
        _clone_rows(tbl, at_index=2, count=extra)  # la plantilla trae UNA fila de alerta
    for i, a in enumerate(alertas):
        r = 2 + i
        _set_cell(tbl.cell(r, 0), _clean(a.get("alerta")), 9)
        _set_cell(tbl.cell(r, 2), _clean(a.get("area")), 9)
        _set_cell(tbl.cell(r, 3), _clean(a.get("unidad_asociada")), 9)
        _set_cell(tbl.cell(r, 4), _clean(a.get("apreciacion")), 9)
    if not alertas:
        for col in (0, 2, 3, 4):
            _set_cell(tbl.cell(2, col), "")
    _set_cell(tbl.cell(5 + extra, 0), _v(c, "inteligencia.analisis_meteo"), 9)
    _set_cell(tbl.cell(8 + extra, 0), _v(c, "inteligencia.proyeccion_meteo_48h"), 9)
    # La fila separadora sobre la banda 5 trae fill dorado horneado (se veía como una
    # barrita amarilla suelta encima del título) -> se pinta blanca.
    for col in range(5):
        celda = tbl.cell(6 + extra, col)
        celda.fill.solid()
        celda.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # --- Reflow: las tablas 2 y 3 suben/bajan según lo que ocupe la anterior ---
    top21 = max(_bottom(t9) + Inches(0.38), Inches(2.35))
    _move(t21, top=int(top21))
    top10 = _bottom(t21) + Inches(0.38)
    _move(t10, top=int(top10))

    _set_text(s2["object 17"].text_frame, fstr)
    _shape_fit(s2["object 17"], pt=15.5, wrap=False)


# ---------------------------------------------------------------------------
#  OPERACIONES — tablas (slide idx3)
# ---------------------------------------------------------------------------
def _s3_operaciones(slide, c: dict, fstr: str) -> None:
    s3 = _by_name(slide)

    # El título 'OPERACIONES EN CURSO' es un PLACEHOLDER de 3.80": LibreOffice ignora
    # wrap=none en placeholders y partía 'CURSO' a la 2ª línea -> caja ensanchada para
    # que el texto (28pt ~4.0") quepa aunque el renderer envuelva.
    _move(s3["object 7"], width=Inches(4.40))

    t15, t13, t14 = s3["object 15"], s3["object 13"], s3["object 14"]
    t12, t16 = s3["object 12"], s3["object 16"]

    # 1. OPERACIONES MILITARES EN CURSO (CATÁSTROFES)
    _fill_table_list(t15, _items(c, "operaciones.ops_catastrofe"), start_row=2,
                     cols={0: "cantidad", 2: "tipo", 3: "unidad_origen", 4: "zona_empleo",
                           5: "inicio_empleo"}, size=9.5, font=FGM, underline=False)

    # 2/3. RELEVOS MACROZONA NORTE / SUR — N.º y FECHA vienen fusionados verticalmente
    # en la plantilla (un relevo con varias unidades); cada relevo del briefing es una
    # fila completa, así que se deshace el merge para que todos muestren su N.º/fecha.
    _unmerge_col(t13, col=0, rows=[2, 3, 4])
    _unmerge_col(t13, col=1, rows=[2, 3, 4])
    _unmerge_col(t14, col=0, rows=[2, 3, 4, 5])
    _unmerge_col(t14, col=1, rows=[2, 3, 4, 5])
    _fill_table_list(t13, _items(c, "operaciones.relevos_norte"),
                     start_row=2, cols={0: "n", 1: "fecha", 2: "jaf", 3: "unidades"},
                     size=9.5, font=FGM, underline=False,
                     aligns={0: PP_ALIGN.CENTER, 1: PP_ALIGN.CENTER})
    _fill_table_list(t14, _items(c, "operaciones.relevos_sur"),
                     start_row=2, cols={0: "n", 1: "fecha", 2: "ft", 3: "unidades"},
                     size=9.5, font=FGM, underline=False,
                     aligns={0: PP_ALIGN.CENTER, 1: PP_ALIGN.CENTER})

    # 5. OPS ANTÁRTICA — cabeceras a 10pt para que 'DOTACIÓN ESTIVAL' envuelva sin partirse.
    _cell_fit(t12, {(1, 0): 10, (1, 1): 10, (1, 2): 10}, wrap=True)
    _fill_table_list(t12, _items(c, "operaciones.ops_antartica"),
                     start_row=2, cols={0: "bae", 1: "dotacion_estival", 2: "cpccgu"},
                     size=9.5, font=FGM, underline=False)

    # 6. OPS EXTRANJERO — columna OPERACIÓN más ancha + cabeceras 10pt.
    _set_col_widths(t16, {0: Inches(1.12), 1: Inches(1.71)})
    _cell_fit(t16, {(1, 0): 10, (1, 1): 10, (1, 2): 10, (1, 3): 10}, wrap=True)
    _fill_table_list(t16, _items(c, "operaciones.ops_extranjero"),
                     start_row=2, cols={0: "operacion", 1: "efectivos", 2: "ubicacion",
                                        3: "repliegue"}, size=9.5, font=FGM, underline=False)

    # --- Reflow vertical: bloque 1 arriba; 2/3 y 5/6 se reparten el espacio libre ---
    top1 = Inches(1.18)
    h1 = _est_table_h(t15)
    h23 = max(_est_table_h(t13), _est_table_h(t14))
    h56 = max(_est_table_h(t12), _est_table_h(t16))
    libre = (Inches(10.55) - top1) - (h1 + h23 + h56)
    gap = max(Inches(0.35), min(Inches(1.15), int(libre // 3)))
    top23 = top1 + h1 + gap
    _move(t13, top=int(top23))
    _move(t14, top=int(top23))
    top56 = top23 + h23 + gap
    _move(t12, top=int(top56))
    _move(t16, top=int(top56))

    _set_text(s3["object 8"].text_frame, fstr)
    _shape_fit(s3["object 8"], pt=15.5, wrap=False)


# ---------------------------------------------------------------------------
#  OPERACIONES — KPIs + unidades en terreno (slide idx4)
# ---------------------------------------------------------------------------
def _s4_operaciones_kpi(slide, c: dict, fstr: str) -> None:
    s4 = _by_name(slide)
    kpi = "operaciones.unidades_terreno_kpi"

    # KPIs: rótulos Franklin Gothic Medium 12 (template) y cifras al tamaño horneado
    # (Trebuchet 28 / 23.5): solo se reemplaza el texto.
    _set_para(s4["object 23"].text_frame, 1, _v(c, f"{kpi}.fuerza_en_terreno"))
    _set_para(s4["object 39"].text_frame, 2, _v(c, f"{kpi}.unidades_en_terreno"))
    _set_para(s4["object 46"].text_frame, 1, _v(c, f"{kpi}.lugares_activos"))
    for name in ("object 23", "object 39", "object 46"):
        tf = s4[name].text_frame
        tf.word_wrap = True
        try:
            tf.auto_size = MSO_AUTO_SIZE.NONE
        except Exception:
            pass

    # 'Fuerza en terreno': la cifra del template trae interlineado FIJO de 33.5pt
    # (rótulo en 2 líneas + 33.5 > alto de la caja: la cifra caía fuera del recuadro
    # azul) y el rótulo una sangría colgante de 0.47" que en LibreOffice lo partía en
    # 3 líneas. Se quitan lnSpc fijo y sangrías, todo centrado, cifra a 26pt: cabe.
    tf23 = s4["object 23"].text_frame
    for p in tf23.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        pPr = p._p.find(qn("a:pPr"))
        if pPr is not None:
            for attr in ("marL", "marR", "indent"):
                if pPr.get(attr):
                    del pPr.attrib[attr]
            lnSpc = pPr.find(qn("a:lnSpc"))
            if lnSpc is not None:
                pPr.remove(lnSpc)
    for r in tf23.paragraphs[1].runs:
        r.font.size = Pt(26)

    # 'Actividad principal': rótulo FGM 12 + valor en línea aparte a 8pt. La caja llega
    # justo hasta la divisoria del icono (x=3.21) y el valor pierde el tracking negativo
    # heredado del rótulo (a 8pt lo volvía ilegible y desbordaba sobre la divisoria).
    act = s4["object 51"]
    _move(act, left=Inches(2.19), width=Inches(1.00))
    tf = act.text_frame
    n_label = len(tf.paragraphs)
    _append_line(tf, _v(c, f"{kpi}.actividad_principal"))
    _fit(tf, None, wrap=True)
    for i in range(n_label):
        _para_fit(tf, i, 12)
    for r in tf.paragraphs[-1].runs:
        r.font.size = Pt(8)
        rPr = r._r.get_or_add_rPr()
        if rPr.get("spc"):
            del rPr.attrib["spc"]

    # 7. UNIDADES EN TERRENO: se deshace el merge de COMANDO MATRIZ (cada fila su comando).
    t37 = s4["object 37"]
    _unmerge_col(t37, col=1, rows=[2, 3, 4])
    _cell_fit(t37, {(1, i): 10.5 for i in range(6)}, wrap=True)
    _fill_table_list(t37, _items(c, "operaciones.unidades_terreno"), start_row=2,
                     cols={0: "actividad", 1: "comando_matriz", 2: "ur", 3: "fechas",
                           4: "ubicacion", 5: "fuerza"}, size=9)

    _set_text(s4["object 33"].text_frame, fstr)
    _shape_fit(s4["object 33"], pt=15.5, wrap=False)


# ---------------------------------------------------------------------------
#  LOGÍSTICA (slide idx5): TOTAL/NOP reales sobre la imagen horneada
# ---------------------------------------------------------------------------
def _pic_por_tamano(grupo, min_w_in: float, min_h_in: float):
    for sub in grupo.shapes:
        if sub.shape_type == MSO_SHAPE_TYPE.PICTURE and sub.width >= Inches(min_w_in) \
                and sub.height >= Inches(min_h_in):
            return sub
    return None


def _swap_pic_blob(slide, pic, png: bytes) -> None:
    blip = pic._element.blip_rId
    part = slide.part.related_part(blip)
    part._blob = png


def _s5_logistica(slide, c: dict, fstr: str) -> None:
    s5 = _by_name(slide)
    _set_text(s5["object 9"].text_frame, fstr)
    _shape_fit(s5["object 9"], pt=15.5, wrap=False)

    # TOTAL / NOP: la columna izquierda es una IMAGEN con '389'/'69' de ejemplo horneados.
    # Se blanquean esas casillas en el PNG embebido y se dibujan los valores reales.
    res = _dget(c, "logistica.resumen")
    items = _items(c, "logistica.operacionalidad")
    tot = _clean(res.get("total")) or (_miles(sum(_num(i.get("total")) or 0 for i in items))
                                       if items else "")
    nop = _clean(res.get("nop")) or (_miles(sum(_num(i.get("nop")) or 0 for i in items))
                                     if items else "")
    grupo = s5.get("object 10")
    pic = _pic_por_tamano(grupo, 1.8, 5.0) if grupo is not None else None
    if pic is not None:
        try:
            im = PILImage.open(io.BytesIO(pic.image.blob)).convert("RGB")
            sx, sy = im.width / 173.0, im.height / 563.0
            d = ImageDraw.Draw(im)

            def celda(x0, y0, x1, y1, color, texto):
                caja = (x0 * sx, y0 * sy, x1 * sx, y1 * sy)
                d.rectangle(caja, fill=color)
                if not texto:
                    return
                px = int(30 * sy)
                font = _pil_font("Times New Roman", True, px)
                while font is not None and px > 8 and \
                        d.textlength(texto, font=font) > (caja[2] - caja[0]) - 10:
                    px -= 1
                    font = _pil_font("Times New Roman", True, px)
                if font is None:
                    return
                bb = d.textbbox((0, 0), texto, font=font)
                x = caja[0] + ((caja[2] - caja[0]) - (bb[2] - bb[0])) / 2 - bb[0]
                y = caja[1] + ((caja[3] - caja[1]) - (bb[3] - bb[1])) / 2 - bb[1]
                d.text((x, y), texto, font=font, fill=(0, 0, 0))

            celda(4, 86, 96, 149, (255, 255, 255), tot)
            celda(103, 86, 169, 149, (255, 0, 0), nop)
            buf = io.BytesIO()
            im.save(buf, format="PNG")
            _swap_pic_blob(slide, pic, buf.getvalue())
        except Exception:
            pass  # si la imagen embebida cambia, se deja como está (no romper el export)


# ---------------------------------------------------------------------------
#  Relleno de las 6 páginas
# ---------------------------------------------------------------------------
def _rellenar(slides, c: dict, fecha: datetime) -> None:
    fstr = _fecha_es(fecha)

    _s0_portada(slides[0], c)
    _s1_personal(slides[1], c, fecha)
    _s2_inteligencia(slides[2], c, fstr)
    _s3_operaciones(slides[3], c, fstr)
    _s4_operaciones_kpi(slides[4], c, fstr)
    _s5_logistica(slides[5], c, fstr)

    # Cabeceras comunes (idx1..idx5): caja negra de sección a 20pt, título "SITUACIÓN ..."
    # al tamaño de la plantilla (Trebuchet MS 28, una línea) y fecha 15.5. Con las fuentes
    # reales instaladas caben en una línea; wrap=False evita que LibreOffice las parta.
    cabeceras = {
        1: ("object 18", "object 19"),
        2: ("object 15", "object 16"),
        3: ("object 6", "object 7"),
        4: ("object 31", "object 32"),
        5: ("object 7", "object 8"),
    }
    for idx, (negra, blanco) in cabeceras.items():
        nm = _by_name(slides[idx])
        _shape_fit(nm.get(negra), pt=20, wrap=False)
        _shape_fit(nm.get(blanco), pt=28, wrap=False)


def _graficos(prs, slides, c: dict) -> None:
    # --- PORTADA: mini-dona dentro de PARTE DE FUERZA INSTITUCIONAL ---
    pf = _dget(c, "personal.parte_fuerza_institucional")
    mini = _chart_mini_donut(pf)
    _reemplazar(slides[0], ["object 35"], mini,
                rect=(Inches(2.31), Inches(1.86), Inches(0.97), Inches(0.97)))

    # --- PÁG. 2: dona grande (se quitan también los cuadraditos/rótulos de leyenda
    #     y el 'TOTAL' horneados; la dona nueva trae sus propios rótulos) ---
    donut = _chart_donut(pf)
    _reemplazar(slides[1], ["object 3", "object 7", "object 8", "object 9", "object 10",
                            "object 11"], donut,
                rect=(Inches(0.33), Inches(1.62), Inches(2.75), Inches(2.85)))

    # --- PÁG. 2: barras por macrozona, ancladas entre las bandas 2 y 3 ---
    nm1 = _by_name(slides[1])
    banda2, banda3 = nm1.get("object 13"), nm1.get("object 24")
    m_top = (banda2.top + banda2.height + Inches(0.12)) if banda2 else Inches(5.70)
    m_bot = (banda3.top - Inches(0.10)) if banda3 else Inches(8.20)
    m_left = Inches(3.58)
    barras = _chart_macrozonas(_get(c, "personal.macrozonas") or [])
    _reemplazar(slides[1], _MACRO_NAMES, barras,
                rect=(m_left, m_top, prs.slide_width - Inches(0.28) - m_left, m_bot - m_top))

    # --- PÁG. 2: barras de fuerza de catástrofe (gráfico que faltaba), bajo la banda 3 ---
    c_top = (banda3.top + banda3.height + Inches(0.12)) if banda3 else Inches(8.86)
    c_h = min(Inches(1.75), prs.slide_height - Inches(0.25) - c_top)
    cata = _chart_catastrofe(_dget(c, "personal.catastrofe"))
    _reemplazar(slides[1], _CATA_NAMES, cata,
                rect=(m_left, c_top, prs.slide_width - Inches(0.28) - m_left, c_h))

    # --- PÁG. 5: Traslados por UAC (tarjeta regenerada con matplotlib) ---
    uac = _chart_uac(_get(c, "operaciones.traslados_uac") or [])
    s4 = _by_name(slides[4])
    titulo_uac = s4.get("object 20")
    if uac:
        _reemplazar(slides[4], _UAC_NAMES, uac,
                    rect=(Inches(4.25), Inches(1.04), Inches(4.10), Inches(2.09)))
        if titulo_uac is not None:
            _to_front(titulo_uac)  # el título (Calibri 18.5) queda sobre la tarjeta nueva
    else:
        # Sin datos: se eliminan las cifras de ejemplo horneadas (nunca inventar).
        _reemplazar(slides[4], ["object 7", "object 8", "object 9", "object 10", "object 11"],
                    None)

    # --- PÁG. 6: operacionalidad logística, bajo la fila de cabeceras de la tabla ---
    oper = _chart_operacionalidad(_get(c, "logistica.operacionalidad") or [],
                                  _dget(c, "logistica.resumen"))
    if oper:
        nm = _by_name(slides[5])
        grupo = nm.get("object 10")
        lado = _pic_por_tamano(grupo, 1.8, 5.0) if grupo is not None else None
        left = (lado.left + lado.width + Inches(0.18)) if lado is not None else Inches(2.66)
        top = Inches(3.05)  # bajo la fila CAMPAÑA/COMBATE/... (que termina ~2.6")
        width = prs.slide_width - Inches(0.30) - left
        height = int(width * 5.2 / 7.4)  # conserva el aspecto del gráfico (7.4 x 5.2)
        _reemplazar(slides[5], ["object 13"], oper, rect=(left, top, width, height))


# ---------------------------------------------------------------------------
#  PPTX -> PDF (LibreOffice headless) — python-pptx no exporta PDF
# ---------------------------------------------------------------------------
def _pptx_a_pdf(data: bytes) -> bytes:
    import glob
    import subprocess
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "reporte.pptx")
        with open(src, "wb") as fh:
            fh.write(data)
        env = {**os.environ, "HOME": tempfile.gettempdir()}
        try:
            subprocess.run(
                [settings.soffice_bin, "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                check=True, env=env, capture_output=True, timeout=180,
            )
        except FileNotFoundError as e:
            raise RuntimeError(
                f"LibreOffice ('{settings.soffice_bin}') no está instalado; se requiere para "
                "exportar el briefing a PDF (RF-008). Ajusta SOFFICE_BIN."
            ) from e
        except subprocess.CalledProcessError as e:
            raise RuntimeError(
                "LibreOffice falló al convertir a PDF: " + e.stderr.decode("utf-8", "ignore")[:500]
            ) from e
        pdfs = glob.glob(os.path.join(tmp, "*.pdf"))
        if not pdfs:
            raise RuntimeError("LibreOffice no generó el PDF esperado.")
        with open(pdfs[0], "rb") as fh:
            return fh.read()


def generar_pdf(contenido: dict[str, Any], titulo: str, fecha: datetime) -> bytes:
    """Rellena la plantilla institucional y devuelve el PDF final (RF-008)."""
    prs = Presentation(_TEMPLATE)
    slides = list(prs.slides)
    _rellenar(slides, contenido, fecha)
    _graficos(prs, slides, contenido)
    buf = io.BytesIO()
    prs.save(buf)
    return _pptx_a_pdf(buf.getvalue())
