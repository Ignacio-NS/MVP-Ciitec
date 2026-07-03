"""
Filtro de salida — EXPORTACIÓN (RF-008): briefing -> PDF / Word / texto plano.

- PDF (maestro): se RELLENA la plantilla institucional de 6 páginas (.pptx) con los datos
  del briefing y se convierte a PDF con LibreOffice headless. Ver `reporte_pptx.py`.
- Word: documento .docx generado con python-docx (mismo contenido por secciones).
- Texto: volcado estructurado.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from . import reporte_pptx


def render_pdf(contenido: dict[str, Any], titulo: str, fecha: datetime) -> bytes:
    """PDF institucional: plantilla .pptx de 6 páginas rellenada -> LibreOffice -> PDF."""
    return reporte_pptx.generar_pdf(contenido, titulo, fecha)


def render_texto(contenido: dict[str, Any], titulo: str, fecha: datetime) -> bytes:
    L = [f"{titulo}  ({fecha.strftime('%d-%m-%Y')})", "=" * 60, "", "RESUMEN EJECUTIVO:"]
    for b in contenido.get("resumen_ejecutivo", []) or ["-.-"]:
        L.append(f"  • {b}")
    sit = contenido.get("situacion") or {}
    L += ["", "SITUACIÓN:", "  " + (sit.get("resumen") or "-.-")]
    for a in sit.get("aspectos", []) or []:
        L.append(f"  • {a}")
    L += ["", "ASUNTOS CRÍTICOS:"]
    for a in contenido.get("asuntos_criticos", []) or []:
        L.append(f"  - {a.get('asunto','-.-')} | {a.get('impacto','-.-')} | {a.get('responsable','-.-')}")
    L += ["", "PROYECCIÓN 24-72H:", "  " + (contenido.get("proyeccion_24_72h") or "-.-")]
    return "\n".join(L).encode("utf-8")


def render_word(contenido: dict[str, Any], titulo: str, fecha: datetime) -> bytes:
    import docx

    doc = docx.Document()
    doc.add_heading(titulo, level=0)
    doc.add_paragraph(fecha.strftime("%d-%m-%Y"))

    doc.add_heading("Resumen ejecutivo", level=1)
    for b in contenido.get("resumen_ejecutivo", []) or ["-.-"]:
        doc.add_paragraph(b, style="List Bullet")

    sit = contenido.get("situacion") or {}
    doc.add_heading("Situación", level=1)
    doc.add_paragraph(sit.get("resumen") or "-.-")
    for a in sit.get("aspectos", []) or []:
        doc.add_paragraph(a, style="List Bullet")

    doc.add_heading("Asuntos críticos", level=1)
    ac = contenido.get("asuntos_criticos", []) or []
    if ac:
        t = doc.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Asunto", "Impacto", "Responsable"
        for a in ac:
            cel = t.add_row().cells
            cel[0].text = a.get("asunto", "-.-")
            cel[1].text = a.get("impacto", "-.-")
            cel[2].text = a.get("responsable", "-.-")

    doc.add_heading("Proyección 24-72h", level=1)
    doc.add_paragraph(contenido.get("proyeccion_24_72h") or "-.-")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render(contenido: dict[str, Any], formato: str, titulo: str, fecha: datetime) -> tuple[bytes, str, str]:
    """Devuelve (bytes, content_type, extension) según el formato."""
    formato = formato.upper()
    if formato == "PDF":
        return render_pdf(contenido, titulo, fecha), "application/pdf", "pdf"
    if formato == "WORD":
        return (
            render_word(contenido, titulo, fecha),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )
    return render_texto(contenido, titulo, fecha), "text/plain; charset=utf-8", "txt"
